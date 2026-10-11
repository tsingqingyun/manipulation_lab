"""Build the editable English assignment deck and its Chinese presenter script.

Authoring-only dependencies: python-pptx, matplotlib, numpy, Pillow.
The lab runtime and original results are not modified.
"""
import csv
import argparse
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
A = HERE / "assets"
CONTENT = json.loads((HERE / "slides.json").read_text())
NAVY = "101F2C"
TEAL = "087F80"
MINT = "A8E9C7"
PAPER = "F5F4EF"
WHITE = "FFFFFF"
GRAY = "546574"
PALE = "E5EAE5"
ORANGE = "D97542"
FONT = "DejaVu Sans"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = CONTENT["title"]
prs.core_properties.subject = "Introduction to Robotics — Assignment Introduction"
prs.core_properties.author = "Manipulation Learning Lab"
prs.core_properties.keywords = "Panda, PickPlaceCan, behavior cloning, assignment"


def color(s): return RGBColor.from_string(s)


def rect(s, x, y, w, h, fill, radius=False, line=None):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
    else:
        shape.line.fill.background()
    if radius: shape.adjustments[0] = .08
    return shape


def text(s, value, x, y, w, h, size=20, fill=NAVY, bold=False, align=None):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.01)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color(fill)
        p.space_after = Pt(5)
        if align is not None: p.alignment = align
    return box


def label(s, value, x, y, w=4, fill=TEAL):
    return text(s, value, x, y, w, .25, 10, fill, True)


def line(s, x1, y1, x2, y2, fill=GRAY, width=1.5):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = color(fill); sh.line.width = Pt(width)
    return sh


def arrow(s, x, y, w=.35, h=.22, fill=TEAL):
    sh = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color(fill); sh.line.fill.background()
    return sh


def picture(s, name, x, y, w, h=None):
    return s.shapes.add_picture(str(A / name), Inches(x), Inches(y), width=Inches(w),
                                height=None if h is None else Inches(h))


def chip(s, value, x, y, w, fill=PALE, fg=NAVY, size=12):
    rect(s, x, y, w, .4, fill, True)
    text(s, value, x+.1, y+.055, w-.2, .29, size, fg, True)


def slide(number, dark=False, section=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = color(NAVY if dark else PAPER)
    foreground = WHITE if dark else NAVY
    label(s, "INTRODUCTION TO ROBOTICS  /  ASSIGNMENT", .55, .35, 8, MINT if dark else TEAL)
    if number != 1:
        text(s, CONTENT["slides"][number-1]["title"], .55, .85, 12.2, .62, 30, foreground, True)
    line(s, .55, 7.08, 12.78, 7.08, "314350" if dark else "D8DED8", .8)
    text(s, "MANIPULATION LEARNING LAB", .55, 7.2, 5, .2, 9, "BAC9CF" if dark else GRAY)
    text(s, section, 5, 7.2, 6.95, .2, 9, "BAC9CF" if dark else GRAY, align=PP_ALIGN.RIGHT)
    text(s, f"{number:02d} / 09", 12, 7.18, .78, .25, 10, MINT if dark else TEAL, align=PP_ALIGN.RIGHT)
    item = CONTENT["slides"][number-1]
    notes = f"第 {number} 页｜{item['seconds']} 秒\n{item['notes_zh']}\n\n来源核对：\n" + "\n".join(item["sources"])
    s.notes_slide.notes_text_frame.text = notes
    return s


def make_figures():
    plt.rcParams.update({"font.family": FONT, "font.size": 12, "axes.labelcolor": '#'+NAVY,
        "text.color": '#'+NAVY, "xtick.color": '#'+GRAY, "ytick.color": '#'+GRAY,
        "axes.edgecolor": '#'+PALE, "axes.spines.top": False, "axes.spines.right": False,
        "figure.facecolor": '#'+WHITE, "axes.facecolor": '#'+WHITE})
    fig, ax = plt.subplots(figsize=(6.1, 3.65))
    for variant, co, name in [("lr_a", TEAL, "0.001"), ("lr_b", ORANGE, "0.0003")]:
        rows = list(csv.DictReader((A / f"{variant}_training_log.csv").open()))
        epochs = [int(r["epoch"]) for r in rows]
        ax.plot(epochs, [float(r["training_loss"]) for r in rows], color='#'+co, alpha=.38, lw=2)
        val = [float(r["validation_loss"]) for r in rows]
        ax.plot(epochs, val, color='#'+co, lw=2.3, label=f"Validation · lr {name}")
        idx = int(np.argmin(val)); ax.axvline(epochs[idx], color='#'+co, ls=':', lw=1.4)
        ax.scatter([epochs[idx]], [val[idx]], color='#'+co, zorder=5, s=40)
    ax.set(xlabel="Epoch (one pass through training data)", ylabel="Mean squared action error", yscale="log", xlim=(1,50))
    ax.legend(frameon=False, fontsize=10, loc='upper right')
    ax.grid(axis='y',alpha=.15); fig.tight_layout(); fig.savefig(A / "loss_curves.png", dpi=190); plt.close(fig)
    points = list(csv.DictReader((A / "training_initial_positions.csv").open()))
    result = json.loads((A / "lr_a_evaluation_summary.json").read_text())
    fig, ax = plt.subplots(figsize=(5.8, 4.25))
    ax.scatter([float(p['x']) for p in points], [float(p['y']) for p in points], s=18, color='#ADB9B8', alpha=.55, label='160 training starts')
    for i in range(5):
        pair = [next(e for e in result['episodes'] if e['episode_id']==f'{c}_{i:02d}') for c in ['standard','shifted']]
        xy = [json.loads(e['initial_object_position'])[:2] for e in pair]
        ax.plot([p[0] for p in xy], [p[1] for p in xy], color='#687987', lw=1, zorder=2)
    for condition, co, marker, name in [('standard',TEAL,'o','Standard'),('shifted',ORANGE,'D','Shifted +0.04 m X')]:
        xy=np.array([json.loads(e['initial_object_position'])[:2] for e in result['episodes'] if e['condition']==condition])
        ax.scatter(xy[:,0],xy[:,1],s=80,color='#'+co,marker=marker,edgecolor='white',zorder=4,label=name)
    ax.set(xlabel='Can initial world X (m)',ylabel='Can initial world Y (m)')
    ax.legend(frameon=False,fontsize=10,loc='upper left'); ax.grid(alpha=.12)
    fig.tight_layout();fig.savefig(A / 'positions.png',dpi=190);plt.close(fig)
    trace=np.load(A / 'failure_trace.npz');t=np.arange(len(trace['object_position']))/20
    fig,axs=plt.subplots(2,1,figsize=(6.2,3.3),sharex=True)
    axs[0].plot(t,trace['object_position'][:,2],color='#'+TEAL,lw=2)
    axs[0].set_ylabel('Can height (m)')
    axs[1].plot(t,trace['gripper_width'],color='#'+ORANGE,lw=2)
    axs[1].set(ylabel='Gripper width (m)',xlabel='Simulation time (s)',xlim=(0,25))
    for ax in axs:
        for step in [1,41,81]: ax.axvline(step/20,ls=':',color='#889798',lw=1)
        ax.grid(axis='y',alpha=.15)
    fig.tight_layout();fig.savefig(A / 'failure_trace.png',dpi=190);plt.close(fig)


def make_deck(static_output=None):
    s = slide(1, True, "FROM DEMONSTRATIONS TO ROBOT POLICIES")
    picture(s, 'success_step_089.png', 6.93, .9, 5.85, 5.85)
    rect(s, .55, 1.38, .09, .6, MINT)
    text(s, 'Manipulation\nLearning Lab', .82, 1.27, 5.8, 1.9, 41, WHITE, True)
    text(s, 'From Demonstrations\nto Robot Policies', .82, 3.36, 5.5, 1.02, 25, MINT)
    text(s, CONTENT['project_subtitle'], .82, 4.66, 5.4, .76, 17, 'BAC9CF')
    chip(s, '01  Watch', .82, 5.85, 1.65, '243A47', MINT)
    chip(s, '02  Learn', 2.65, 5.85, 1.65, '243A47', MINT)
    chip(s, '03  Try', 4.48, 5.85, 1.65, '243A47', MINT)
    label(s, 'PANDA  /  PICKPLACECAN  /  REAL SIMULATION', 7.12, 6.43, 5.5, NAVY)

    s = slide(2, section='WATCH THE TASK')
    text(s, 'Can a robot learn to pick and place objects?', .57, 1.43, 12, .37, 18, GRAY)
    s.shapes.add_movie(str(A/'success_half_speed.mp4'), Inches(.55), Inches(1.97), Inches(4.75), Inches(4.75),
                       poster_frame_image=str(A/'success_poster.png'), mime_type='video/mp4')
    chip(s, 'CLICK TO PLAY  ·  15 s  ·  0.5× SPEED', .72, 6.2, 4.35, NAVY, WHITE, 11)
    stages = [('Reach',1),('Grasp',41),('Transport',89),('Place',145)]
    for i,(name,step) in enumerate(stages):
        x=5.7+i*1.76
        picture(s,f'success_step_{step:03d}.png',x,2.06,1.58,1.58)
        text(s,name,x,3.77,1.64,.33,15,NAVY,True)
        if i<3: arrow(s,x+1.6,2.75,.14,.13)
    label(s,'ROBOT',5.72,4.4,1.5);text(s,'Panda arm',5.72,4.7,2,.4,17)
    label(s,'TARGET OBJECT',8.05,4.4,2);text(s,'Red can',8.05,4.7,2,.4,17)
    label(s,'GOAL REGION',10.25,4.4,2.5);text(s,'Can-marked bin',10.25,4.7,2.5,.4,17)
    rect(s,5.7,5.34,7.08,1.37,PALE,True)
    text(s,'Success = inside + released + stable',5.92,5.54,6.66,.4,20,TEAL,True)
    text(s,'Target margin · open gripper · no contact\nStable for 1 s: displacement ≤1 cm; speed ≤5 cm/s',5.92,6.02,6.58,.57,14)

    s = slide(3,section='THE LEARNING WORKFLOW')
    cards=[(.55,'01','Demonstration','A person shows the task.','Record current state\nand the next action.'),
           (4.8,'02','Training','Learn a state → action rule.','Compare with the human action.\nUpdate the model weights.'),
           (9.05,'03','Inference','Let the policy act.','Observe → act → observe again.\nRepeat 20 times per second.')]
    for x,n,title,sub,body in cards:
        rect(s,x,2.03,3.73,3.09,WHITE,True)
        label(s,n,x+.23,2.28,1);text(s,title,x+.23,2.74,3.3,.46,23,NAVY,True)
        text(s,sub,x+.23,3.34,3.23,.69,17,TEAL,True)
        text(s,body,x+.23,4.16,3.23,.78,16,GRAY)
    arrow(s,4.4,3.4,.29,.25);arrow(s,8.65,3.4,.29,.25)
    label(s,'200 HUMAN DEMONSTRATIONS  ·  SPLIT BY COMPLETE EPISODE',.58,5.56,10)
    segments=[(.55,7.68,TEAL,'160 train'),(8.23,1.54,ORANGE,'20 validate'),(9.77,1.54,NAVY,'20 test')]
    for x,w,co,v in segments:
        rect(s,x,5.97,w,.58,co)
        text(s,v,x+.1,6.1,w-.2,.3,13,WHITE,True,PP_ALIGN.CENTER)
    text(s,'Validation selects saved weights. Test results never select the model.',.57,6.68,12,.3,14,GRAY)

    s = slide(4,section='OBSERVATION → POLICY → ACTION')
    rect(s,.55,1.95,4.0,4.55,WHITE,True);rect(s,8.0,1.95,4.78,4.55,WHITE,True)
    label(s,'OBSERVATION  /  CURRENT STATE',.8,2.2,3.8)
    text(s,'16 numbers',.8,2.66,3.45,.66,30,NAVY,True)
    rows=[('3 + 4','Gripper-tip position + orientation'),('2','Two finger joint positions'),('3 + 4','Can position + orientation')]
    for i,(n,t) in enumerate(rows):
        y=3.6+i*.77;text(s,n,.8,y,.9,.4,19,TEAL,True);text(s,t,1.73,y,2.52,.6,15)
    text(s,'World frame · Z points up\nPositions in meters',.8,5.93,3.45,.47,13,GRAY)
    rect(s,5.02,3.18,2.5,1.57,NAVY,True)
    text(s,'POLICY',5.2,3.5,2.15,.43,24,MINT,True,PP_ALIGN.CENTER)
    text(s,'A learned function',5.18,4.06,2.17,.38,13,WHITE,align=PP_ALIGN.CENTER)
    arrow(s,4.65,3.86,.26,.24);arrow(s,7.62,3.86,.26,.24)
    label(s,'ACTION  /  NEXT COMMAND',8.25,2.2,4.1)
    text(s,'7 numbers',8.25,2.66,4.1,.66,30,NAVY,True)
    for i,(n,t) in enumerate([('3','Move in X, Y, Z'),('3','Rotate the gripper'),('1','Close (+) or open (−)')]):
        y=3.6+i*.77;text(s,n,8.25,y,.65,.4,19,TEAL,True);text(s,t,9.0,y,3.5,.47,17)
    text(s,'Commands normalized to [−1, 1]\nPer axis: ±0.05 m; ±0.5 rad (axis-angle)',8.25,5.93,4.22,.47,12,GRAY)
    text(s,'The policy reads state values; the videos help YOU see what happens.',.57,6.7,12,.3,15,TEAL,True)

    s = slide(5,section='BEHAVIOR CLONING, WITHOUT THE MATH DETOUR')
    text(s,'Predict the human action. Measure the error. Improve the prediction.',.57,1.54,12,.5,18,GRAY)
    label(s,'THE STUDENT MODEL',.57,2.31,5)
    chip(s,'16 inputs',.57,2.86,1.48,PALE,NAVY,13)
    arrow(s,2.14,2.97,.26,.17)
    for i in range(3):
        rect(s,2.51+i*.7,2.64,.55,1.31,TEAL,True)
        text(s,'256',2.56+i*.7,3.12,.45,.3,11,WHITE,True,PP_ALIGN.CENTER)
    arrow(s,4.59,2.97,.26,.17)
    chip(s,'7 actions',4.95,2.86,1.35,NAVY,WHITE,12)
    text(s,'Three hidden layers, 256 units each\nA feedforward neural network (MLP)',.58,4.17,5.68,.76,18)
    text(s,'Loss = mean squared difference\nbetween predicted and human actions',.58,5.21,5.7,.7,17,TEAL,True)
    text(s,'Learn from training data.\nSelect saved weights using validation loss.',.58,6.13,5.7,.65,14,GRAY)
    rect(s,6.66,2.28,6.12,4.45,WHITE,True)
    picture(s,'loss_curves.png',6.73,2.54,5.97,3.57)
    text(s,'Real example · pale lines: training\nDotted lines: validation-selected epochs (7 and 11)',6.91,6.16,5.59,.5,12,GRAY)

    s = slide(6,section='YOUR FOUR-PART EXPERIMENT')
    text(s,'Use the prepared course server. Open 01 → 02 → 03 → 04 in Jupyter.',.57,1.55,12,.4,18,GRAY)
    rows=[('01','Explore demonstrations','Play a human demo; inspect inputs and actions.','20 min','Q1'),
          ('02','Train two BC policies','Run both rates; save models, curves and evaluations.','50–60 min','Q2–Q3'),
          ('03','Compare evaluation','Read the SAME saved cases; watch outcomes.','40–50 min','Q4'),
          ('04','Analyze a real failure','Select ≥3 steps; mark the first error; export evidence.','20–30 min','Q5')]
    for i,(n,title,body,duration,q) in enumerate(rows):
        y=2.14+i*.94;rect(s,.55,y,12.23,.8,WHITE,True)
        text(s,n,.78,y+.19,.64,.4,24,TEAL,True)
        text(s,title,1.64,y+.12,4,.37,18,NAVY,True)
        text(s,body,1.64,y+.49,8.45,.29,12,GRAY)
        text(s,duration,10.08,y+.2,1.55,.3,14,GRAY)
        text(s,q,11.85,y+.2,.72,.3,13,TEAL,True,PP_ALIGN.RIGHT)
    rect(s,.55,6.07,12.23,.72,NAVY,True)
    text(s,'ONLY CHANGE THE LEARNING RATE',.78,6.23,4.17,.27,12,MINT,True)
    text(s,'0.001 vs 0.0003  ·  Same data, model, seed 17, 50 epochs, batch 256 & device',5.04,6.23,7.49,.31,11,WHITE)

    s = slide(7,section='MATCHED CASES, HONEST CONCLUSIONS')
    text(s,'2 models × (5 standard + 5 shifted starts) = 20 rollouts',.57,1.53,12,.42,20,TEAL,True)
    rect(s,.55,2.15,5.85,4.26,WHITE,True)
    picture(s,'positions.png',.63,2.18,5.65,4.14)
    label(s,'ONE COMPLETED EXAMPLE  /  STABLE PLACEMENTS',6.81,2.24,5.85)
    columns=[(6.82,2.1,'Learning rate'),(9.27,1.52,'Standard'),(11.1,1.55,'Shifted')]
    for x,w,head in columns: text(s,head,x,2.9,w,.32,14,GRAY,True)
    for i,(rate,a,b) in enumerate([('0.001','2 / 5','3 / 5'),('0.0003','5 / 5','5 / 5')]):
        y=3.48+i*.82
        text(s,rate,6.82,y,2.1,.48,23,NAVY,True)
        text(s,a,9.27,y,1.52,.48,23,TEAL,True)
        text(s,b,11.1,y,1.55,.48,23,ORANGE,True)
    line(s,6.82,5.16,12.68,5.16,'D8DED8',1)
    text(s,'Same source episode + same seed\nShift only the can: +0.04 m along world X',6.82,5.41,5.9,.67,16)
    text(s,'Five cases are a coarse estimate. A position shift does not prove out-of-distribution testing.',.57,6.65,12.16,.36,14,TEAL,True)

    s = slide(8,section='VISUALIZATION → EVIDENCE → EXPLANATION')
    text(s,'Find the first observable error. Separate what you see from why it happened.',.57,1.54,12,.45,18,GRAY)
    for i,step in enumerate([1,41,81]):
        x=.55+i*1.94
        picture(s,f'failure_step_{step:03d}.png',x,2.25,1.76,1.76)
        text(s,f'Step {step}',x,4.17,1.76,.37,16,TEAL,True)
    text(s,'One real failure · primary student model\nAt step 81, the can remains in the source bin.',.57,4.82,5.59,.75,17)
    chip(s,'01  Play video',.55,5.97,1.81,PALE,NAVY,11)
    chip(s,'02  Select steps',2.5,5.97,1.85,PALE,NAVY,11)
    chip(s,'03  Mark error',4.49,5.97,1.81,PALE,NAVY,11)
    rect(s,6.62,2.2,6.16,3.53,WHITE,True)
    picture(s,'failure_trace.png',6.65,2.3,6.1,3.25)
    text(s,'Observation → hypothesis → a test',6.83,6.0,5.64,.43,20,TEAL,True)
    text(s,'Notebook 04 links selected frames to logged state. Check the recorded-frame offset.',.57,6.69,12,.3,14,GRAY)

    s = slide(9,True,section='START WITH NOTEBOOK 01')
    label(s,'YOUR GROUP SUBMITS',.57,1.94,5,MINT)
    text(s,'1–2 page PDF\n+ evidence ZIP',.57,2.46,5.9,1.5,34,WHITE,True)
    text(s,'Answer Q1–Q5 in your own words.\nDownload the ZIP from Notebook 04.\nKeep both models and all outcomes.',.58,4.2,5.58,1.25,18,'BAC9CF')
    text(s,'Generated evidence is an appendix.\nBlank answers or no error marker = DRAFT.',.58,5.88,5.8,.66,14,MINT)
    label(s,'100 POINTS  /  EVIDENCE & REASONING',7.13,1.94,5.5,MINT)
    for i,(v,title) in enumerate([(20,'Demonstration understanding'),(25,'Training & controlled experiment'),(30,'Policy evaluation'),(25,'Failure analysis & discussion')]):
        y=2.58+i*.77
        text(s,str(v),7.13,y,.81,.5,29,MINT,True)
        text(s,title,8.24,y+.07,4.51,.5,17,WHITE)
    rect(s,7.13,6.01,5.65,.7,'243A47',True)
    text(s,'A well-explained failure can earn full credit.',7.33,6.23,5.22,.29,14,MINT,True)

    prs.save(HERE / 'Manipulation_Learning_Lab.pptx')
    if static_output:
        # The PDF is a static handout; export a poster instead of a video annotation.
        # The separately saved teaching PPTX keeps its embedded MP4 unchanged.
        s = prs.slides[1]
        movie = None
        for sh in list(s.shapes):
            if sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}videoFile') is not None:
                movie = sh; break
        assert movie is not None
        index = list(s.shapes._spTree).index(movie._element)
        s.shapes._spTree.remove(movie._element)
        timing = s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}timing')
        if timing is not None: s._element.remove(timing)
        poster = picture(s, 'success_poster.png', .55, 1.97, 4.75, 4.75)
        s.shapes._spTree.remove(poster._element); s.shapes._spTree.insert(index, poster._element)
        for sh in s.shapes:
            if sh.has_text_frame and 'CLICK TO PLAY' in sh.text:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.text = r.text.replace('CLICK TO PLAY', 'VIDEO IN PPTX')
        static_output = Path(static_output)
        static_output.parent.mkdir(parents=True, exist_ok=True)
        prs.save(static_output)
    notes = [f"{CONTENT['title']} — 中文课堂讲解\n总时长：600 秒（10 分钟）\n"]
    for item in CONTENT['slides']:
        notes.append(f"第 {item['number']} 页 · {item['title']} · {item['seconds']} 秒\n\n{item['notes_zh']}\n")
    (HERE / 'speaker_notes_zh.txt').write_text('\n\n'.join(notes))
    print('Built 9 slides; 600 seconds; Chinese notes on every slide.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--static-output', help='Optional poster-only PPTX for PDF export')
    args = parser.parse_args()
    make_figures()
    make_deck(args.static_output)
