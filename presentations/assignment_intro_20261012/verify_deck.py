"""One delivery acceptance: source contracts, PPTX, PDF, media and browser playback."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import subprocess
import zipfile

import imageio.v2 as imageio
import imageio_ffmpeg
import numpy as np
import pymupdf
from PIL import Image
from pptx import Presentation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
A = HERE / 'assets'


def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main(browser):
    data = json.loads((HERE/'slides.json').read_text())
    assert len(data['slides']) == 9
    assert sum(s['seconds'] for s in data['slides']) == data['total_seconds'] == 600
    ppt = Presentation(HERE/'Manipulation_Learning_Lab.pptx')
    assert len(ppt.slides) == 9
    slide_checks, all_text = [], []
    for i,s in enumerate(ppt.slides):
        text = '\n'.join(sh.text for sh in s.shapes if sh.has_text_frame)
        assert not re.search('[\u4e00-\u9fff]', text)
        assert data['slides'][i]['notes_zh'] in s.notes_slide.notes_text_frame.text
        assert re.search('[\u4e00-\u9fff]',s.notes_slide.notes_text_frame.text)
        for sh in s.shapes:
            assert min(sh.left,sh.top) >= 0
            assert sh.left+sh.width <= ppt.slide_width+10 and sh.top+sh.height <= ppt.slide_height+10
        slide_checks.append({'slide':i+1,'visible_language':'English','Chinese_speaker_notes':True,'bounds_passed':True})
        all_text.append(f'SLIDE {i+1}\n{text}')
    (HERE/'visible_slide_text.txt').write_text('\n\n'.join(all_text))
    with zipfile.ZipFile(HERE/'Manipulation_Learning_Lab.pptx') as z:
        names=[n for n in z.namelist() if n.endswith('.mp4')]
        assert len(names)==1
        assert hashlib.sha256(z.read(names[0])).hexdigest()==sha(A/'success_half_speed.mp4')
        embedded_name=names[0]
    video_checks=[]
    for filename in ['success_half_speed.mp4','failure_original.mp4']:
        path=A/filename
        ffmpeg=imageio_ffmpeg.get_ffmpeg_exe()
        subprocess.run([ffmpeg,'-v','error','-i',str(path),'-f','null','-'],check=True,capture_output=True)
        reader=imageio.get_reader(path);meta=reader.get_meta_data();n=sum(1 for _ in reader);reader.close()
        video_checks.append({'file':filename,'duration_seconds':meta['duration'],'fps':meta['fps'],'resolution':meta['size'],'decoded_frames':n})
    assert video_checks[0]['duration_seconds']==15 and video_checks[0]['decoded_frames']==75
    provenance=json.loads((HERE/'media_provenance.json').read_text())
    assert max(provenance['max_abs_state_replay_errors'].values())<=1e-6
    for src in provenance['sources']:
        assert sha(ROOT/src['path'])==src['sha256']
    # A controlled pair must use the same source episodes, seeds and starts.
    summaries=[json.loads((A/f'{v}_evaluation_summary.json').read_text()) for v in ['lr_a','lr_b']]
    for a,b in zip(summaries[0]['episodes'],summaries[1]['episodes']):
        for key in ['source_episode','seed','initial_object_position','condition']:
            assert a[key]==b[key]
    student=json.loads((ROOT/'configs/bc_student.json').read_text())
    design=json.loads((ROOT/'configs/controlled_experiment.json').read_text())
    assert student['hidden_dims']==[256]*3 and student['seed']==17 and student['epochs']==50 and student['batch_size']==256
    assert design['learning_rates']==[.001,.0003] and design['primary_variant']=='lr_a'
    for v,expect in [('lr_a',[2,3]),('lr_b',[5,5])]:
        d=json.loads((A/f'{v}_evaluation_summary.json').read_text())
        counts=[sum(e['stable_success'] for e in d['episodes'] if e['condition']==c) for c in ['standard','shifted']]
        assert counts==expect
    doc=pymupdf.open(HERE/'Manipulation_Learning_Lab.pdf')
    assert len(doc)==9
    (HERE/'preview').mkdir(exist_ok=True)
    sheet=Image.new('RGB',(1920,1080),'#D8DED8')
    pdf_checks=[]
    for i,page in enumerate(doc):
        text=page.get_text()
        assert not re.search('[\u4e00-\u9fff]',text)
        assert len(text)>120
        assert not list(page.annots() or []), 'PDF handout must not contain video annotations'
        pix=page.get_pixmap(matrix=pymupdf.Matrix(1.3,1.3))
        dest=HERE/f'preview/slide_{i+1:02d}.png';pix.save(dest)
        im=Image.open(dest).convert('RGB');im.thumbnail((640,360));sheet.paste(im,((i%3)*640,(i//3)*360))
        # Compare visible nonempty textbox words against the PDF's extracted words.
        expected=' '.join(sh.text for sh in ppt.slides[i].shapes if sh.has_text_frame)
        if i==1: expected=expected.replace('CLICK TO PLAY','VIDEO IN PPTX')
        words=lambda s:set(re.findall(r'[A-Za-z0-9]+',s.lower()))
        missing=words(expected)-words(text)
        assert not missing,(i+1,missing)
        pdf_checks.append({'page':i+1,'all_visible_text_words_present':True,'static_preview':True})
    sheet.save(HERE/'preview/contact_sheet.png')
    browser_results={}
    if browser:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b=p.chromium.launch(executable_path=browser,headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
            page=b.new_page()
            for name in ['success_half_speed.mp4','failure_original.mp4']:
                page.goto((A/name).as_uri())
                page.wait_for_function('document.querySelector("video")?.readyState >= 2')
                page.evaluate('document.querySelector("video").muted = true')
                page.evaluate('document.querySelector("video").currentTime = 0')
                page.evaluate('document.querySelector("video").play()')
                page.wait_for_function('document.querySelector("video").currentTime > 0.5')
                duration=page.evaluate('document.querySelector("video").duration')
                page.evaluate('(t)=>document.querySelector("video").currentTime=t',duration-.5)
                page.wait_for_function('document.querySelector("video").ended')
                browser_results[name]=page.evaluate('(()=>{const v=document.querySelector("video");return {duration:v.duration,width:v.videoWidth,height:v.videoHeight,currentTime:v.currentTime,ended:v.ended,error:v.error};})()')
                assert browser_results[name]['ended'] and browser_results[name]['error'] is None
            browser_results['version']=b.version
            b.close()
    contract_paths=['README.md','docs/assignment_student.md','docs/instructor_guide.md',
                    'configs/environment.yaml','configs/bc_student.json','configs/controlled_experiment.json','configs/evaluation.yaml',
                    'src/manipulation_lab/policy.py','src/manipulation_lab/common.py','src/manipulation_lab/evaluation.py']
    contract_paths += [str(p.relative_to(ROOT)) for p in sorted((ROOT/'notebooks').glob('*.ipynb'))]
    contract={p:sha(ROOT/p) for p in contract_paths}
    (HERE/'source_contract.json').write_text(json.dumps({
        'date':'2026-10-12','base_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),
        'scope':'Current implemented working tree; ongoing course-packaging changes are owned by a separate user session.',
        'files':contract,'rubric':[20,25,30,25],'submission':'1–2 page PDF and Notebook 04 evidence ZIP',
        'control_hz':20,'observation_dimensions':16,'action_dimensions':7,'policy':'deterministic MLP 256/256/256',
        'source_result_role':'One existing completed student example, not a target score or new benchmark.'},indent=2)+'\n')
    report={'passed':True,'slides':slide_checks,'timing_seconds':600,'embedded_video':embedded_name,
            'embedded_video_matches_source_bytes':True,'media':video_checks,'replay_state_max_errors':provenance['max_abs_state_replay_errors'],
            'pdf':pdf_checks,'browser':browser_results,'content_contract_passed':True,
            'limitations':['Native Microsoft PowerPoint playback has not been tested on the instructor laptop.',
                           'PDF is a static preview; play the embedded PPTX media or the separate MP4.',
                           'This presentation does not certify completion of the separate environment-deployment session.']}
    report['deliverable_sha256']={n:sha(HERE/n) for n in ['Manipulation_Learning_Lab.pptx','Manipulation_Learning_Lab.pdf','speaker_notes_zh.txt']}
    (HERE/'verification.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps({'passed':True,'slides':9,'seconds':600,'video_seconds':15,'browser_verified':bool(browser_results)},indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--browser');main(parser.parse_args().browser)
