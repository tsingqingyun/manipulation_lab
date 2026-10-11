"""Build the editable English assignment deck and its Chinese presenter script.

Authoring-only dependencies: python-pptx, matplotlib, numpy, Pillow.
The lab runtime and original results are not modified.
"""
import csv
import argparse
import json
from pptx.oxml.xmlchemy import OxmlElement
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
A = HERE / "assets"
CONTENT = json.loads((HERE / "slides.json").read_text())
NAVY = "20242A"
TEAL = "3E6485"
MINT = "A8E9C7"
PAPER = "FFFFFF"
WHITE = "FFFFFF"
GRAY = "626970"
PALE = "E5EAE5"
ORANGE = "82878C"
FONT = "DejaVu Sans"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = CONTENT["title"]
prs.core_properties.subject = "Introduction to Robotics — Assignment Introduction"
prs.core_properties.author = "Manipulation Learning Lab"
prs.core_properties.keywords = "Panda, PickPlaceCan, behavior cloning, assignment"


def color(s): return RGBColor.from_string(s)


def rect(s, x, y, w, h, fill, radius=False, line=None):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
    else:
        shape.line.fill.background()
    if radius: shape.adjustments[0] = .08
    return shape


def text(s, value, x, y, w, h, size=20, fill=NAVY, bold=False, align=None):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.01)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color(fill)
        p.space_after = Pt(5)
        if align is not None: p.alignment = align
    return box


def label(s, value, x, y, w=4, fill=TEAL):
    return text(s, value, x, y, w, .25, 10, fill, True)


def line(s, x1, y1, x2, y2, fill=GRAY, width=1.5):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = color(fill); sh.line.width = Pt(width)
    style = sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None: sh._element.remove(style)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh


def arrow(s, x, y, w=.35, h=.22, fill=TEAL):
    sh = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color(fill); sh.line.fill.background()
    return sh


def picture(s, name, x, y, w, h=None):
    return s.shapes.add_picture(str(A / name), Inches(x), Inches(y), width=Inches(w),
                                height=None if h is None else Inches(h))


def chip(s, value, x, y, w, fill=PALE, fg=NAVY, size=12):
    rect(s, x, y, w, .4, fill, True)
    text(s, value, x+.1, y+.055, w-.2, .29, size, fg, True)


def slide(number):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = color(WHITE)
    if number != 1:
        text(s, CONTENT['slides'][number-1]['title'], .75, .65, 11.9, .66, 31, NAVY, True)
    text(s, 'Manipulation Learning Lab', .75, 7.15, 6, .22, 10, GRAY)
    text(s, str(number), 12.12, 7.13, .42, .28, 11, GRAY, align=PP_ALIGN.RIGHT)
    item = CONTENT['slides'][number-1]
    s.notes_slide.notes_text_frame.text = (f"第 {number} 页｜{item['seconds']} 秒\n{item['notes_zh']}\n\n来源核对：\n" + "\n".join(item['sources']))
    return s


def equation(s, segments, x, y, w, h=.7, size=29):
    """Editable native text; baseline values are thousandths of a percent."""
    box = text(s, '', x, y, w, h, size)
    box.name = 'Editable equation'
    p = box.text_frame.paragraphs[0]
    for value, baseline, pointsize in segments:
        run = p.add_run(); run.text = value
        run.font.name = FONT; run.font.size = Pt(pointsize or size)
        run.font.color.rgb = color(NAVY)
        if baseline: run._r.get_or_add_rPr().set('baseline', str(52000 if baseline > 0 else baseline))
    return box


def make_figures():
    plt.rcParams.update({"font.family": FONT, "font.size": 12, "axes.labelcolor": '#'+NAVY,
        "text.color": '#'+NAVY, "xtick.color": '#'+GRAY, "ytick.color": '#'+GRAY,
        "axes.edgecolor": '#'+PALE, "axes.spines.top": False, "axes.spines.right": False,
        "figure.facecolor": '#'+WHITE, "axes.facecolor": '#'+WHITE})
    fig, ax = plt.subplots(figsize=(6.1, 3.65))
    for variant, co, name in [("lr_a", TEAL, "0.001"), ("lr_b", ORANGE, "0.0003")]:
        rows = list(csv.DictReader((A / f"{variant}_training_log.csv").open()))
        epochs = [int(r["epoch"]) for r in rows]
        ax.plot(epochs, [float(r["training_loss"]) for r in rows], color='#'+co, alpha=.38, lw=2)
        val = [float(r["validation_loss"]) for r in rows]
        ax.plot(epochs, val, color='#'+co, lw=2.3, label=f"Validation · lr {name}")
        idx = int(np.argmin(val)); ax.axvline(epochs[idx], color='#'+co, ls=':', lw=1.4)
        ax.scatter([epochs[idx]], [val[idx]], color='#'+co, zorder=5, s=40)
    ax.set(xlabel="Epoch (one pass through training data)", ylabel="Mean squared action error", yscale="log", xlim=(1,50))
    ax.legend(frameon=False, fontsize=10, loc='upper right')
    ax.grid(axis='y',alpha=.15); fig.tight_layout(); fig.savefig(A / "loss_curves.png", dpi=190); plt.close(fig)
    points = list(csv.DictReader((A / "training_initial_positions.csv").open()))
    result = json.loads((A / "lr_a_evaluation_summary.json").read_text())
    fig, ax = plt.subplots(figsize=(5.8, 4.25))
    ax.scatter([float(p['x']) for p in points], [float(p['y']) for p in points], s=18, color='#ADB9B8', alpha=.55, label='160 training starts')
    for i in range(5):
        pair = [next(e for e in result['episodes'] if e['episode_id']==f'{c}_{i:02d}') for c in ['standard','shifted']]
        xy = [json.loads(e['initial_object_position'])[:2] for e in pair]
        ax.plot([p[0] for p in xy], [p[1] for p in xy], color='#687987', lw=1, zorder=2)
    for condition, co, marker, name in [('standard',TEAL,'o','Standard'),('shifted',ORANGE,'D','Shifted +0.04 m X')]:
        xy=np.array([json.loads(e['initial_object_position'])[:2] for e in result['episodes'] if e['condition']==condition])
        ax.scatter(xy[:,0],xy[:,1],s=80,color='#'+co,marker=marker,edgecolor='white',zorder=4,label=name)
    ax.set(xlabel='Can initial world X (m)',ylabel='Can initial world Y (m)')
    ax.legend(frameon=False,fontsize=10,loc='upper left'); ax.grid(alpha=.12)
    fig.tight_layout();fig.savefig(A / 'positions.png',dpi=190);plt.close(fig)
    trace=np.load(A / 'failure_trace.npz');t=np.arange(len(trace['object_position']))/20
    fig,axs=plt.subplots(2,1,figsize=(6.2,3.3),sharex=True)
    axs[0].plot(t,trace['object_position'][:,2],color='#'+TEAL,lw=2)
    axs[0].set_ylabel('Can height (m)')
    axs[1].plot(t,trace['gripper_width'],color='#'+ORANGE,lw=2)
    axs[1].set(ylabel='Gripper width (m)',xlabel='Simulation time (s)',xlim=(0,25))
    for ax in axs:
        for step in [1,41,81]: ax.axvline(step/20,ls=':',color='#889798',lw=1)
        ax.grid(axis='y',alpha=.15)
    fig.tight_layout();fig.savefig(A / 'failure_trace.png',dpi=190);plt.close(fig)


def make_deck(static_output=None):
    s = slide(1)
    text(s,'Introduction to Robotics',.78,1.08,5.8,.4,17,GRAY)
    text(s,'Manipulation\nLearning Lab',.75,2.03,6.2,1.62,38,NAVY,True)
    text(s,'From Demonstrations to Robot Policies',.78,4.0,5.6,.86,22,TEAL)
    text(s,'Learning Robot Manipulation\nfrom Demonstrations',.78,5.6,5.65,.73,16,GRAY)
    picture(s,'success_step_089.png',7.06,1.07,5.48,5.48)

    s = slide(2)
    text(s,'Panda arm · red can · can-marked target bin',.78,1.48,12,.38,18,GRAY)
    s.shapes.add_movie(str(A/'success_half_speed.mp4'), Inches(.75), Inches(2.0), Inches(4.65), Inches(4.65),
                      poster_frame_image=str(A/'success_poster.png'), mime_type='video/mp4')
    text(s,'Click to play · 15 s · half speed',.78,6.72,4.6,.32,13,GRAY)
    for i,(name,step) in enumerate([('Reach',1),('Grasp',41),('Transport',89),('Place',145)]):
        x=5.9+i*1.68
        picture(s,f'success_step_{step:03d}.png',x,2.08,1.54,1.54)
        text(s,name,x,3.8,1.63,.36,15,NAVY)
    line(s,5.91,4.49,12.49,4.49,'D6D9DC',.8)
    text(s,'Stable released placement',5.91,4.8,6.5,.43,22,NAVY,True)
    text(s,'Inside the target region, with a radius margin\nGripper open and no longer touching the can\nStable for 1 s after release',5.91,5.45,6.5,1.12,17)
    text(s,'Displacement ≤1 cm · speed ≤5 cm/s',5.91,6.67,6.5,.3,13,GRAY)

    s = slide(3)
    cols=[(.78,'Demonstrations','Recorded observations\nand human actions.'),
          (5.0,'Training','Fit a model to those\nobservation–action pairs.'),
          (9.1,'Inference','Use the current observation\nto choose the next action.')]
    for x,title,body in cols:
        text(s,title,x,2.22,3.5,.55,24,NAVY,True)
        line(s,x,2.97,x+3.42,2.97,'D6D9DC',.8)
        text(s,body,x,3.3,3.43,1.08,19)
    text(s,'→',4.44,2.27,.47,.4,23,GRAY)
    text(s,'→',8.56,2.27,.47,.4,23,GRAY)
    text(s,'Observe → act → observe again, at 20 Hz',.78,4.87,11.8,.51,23,TEAL)
    line(s,.78,5.66,12.49,5.66,'D6D9DC',.8)
    text(s,'200 complete demonstrations',.78,6.03,4.7,.4,19,NAVY,True)
    text(s,'160 train     20 validation     20 test',6.18,6.03,6.3,.4,19)
    text(s,'Select saved model weights using validation loss; reserve test cases for evaluation.',.78,6.62,11.8,.35,14,GRAY)

    s = slide(4)
    text(s,'Observation: 16 values',.78,1.84,5.5,.47,23,NAVY,True)
    text(s,'Action: 7 values',7.31,1.84,5.2,.47,23,NAVY,True)
    left=[('7','Gripper-tip position and orientation'),('2','Finger joint positions'),('7','Can position and orientation')]
    right=[('3','Translation in X, Y, Z'),('3','Rotation of the gripper'),('1','Close (+) or open (−)')]
    for x,rows in [(.78,left),(7.31,right)]:
        line(s,x,2.5,x+5.15,2.5,'D6D9DC',.8)
        for i,(n,desc) in enumerate(rows):
            y=2.78+i*.6
            text(s,n,x,y,.5,.37,20,TEAL)
            text(s,desc,x+.65,y,4.63,.42,18)
    text(s,'World frame · Z up · positions in meters',.78,4.32,5.45,.32,14,GRAY)
    text(s,'Normalized commands in [−1, 1]',7.31,4.32,5.2,.32,14,GRAY)
    equation(s,[('â',0,33),('t',-24000,21),(' = π',0,33),('θ',-24000,21),('(o',0,33),('t',-24000,21),(')',0,33)],.8,4.91,4.0,.83)
    text(s,'o: normalized observation    θ: learned weights\nâ: predicted action    t: current control step',5.07,4.93,7.4,.83,17)
    text(s,'Model: 16 → 256 → 256 → 256 → 7',.8,6.13,7.2,.41,20,NAVY)
    text(s,'Training-only normalization · fixed weights during inference',.8,6.69,11.7,.31,14,GRAY)

    s = slide(5)
    text(s,'Match the demonstrated action, then update the model weights.',.78,1.46,11.8,.4,18,GRAY)
    text(s,'Mean squared action error',.78,2.12,7.1,.4,20,NAVY,True)
    equation(s,[('L',0,30),('B',-22000,18),('(θ) =',0,30)],.78,2.86,1.77,.72)
    text(s,'1',2.59,2.73,.5,.33,21,NAVY,align=PP_ALIGN.CENTER)
    line(s,2.51,3.12,3.2,3.12,NAVY,1)
    text(s,'7B',2.5,3.22,.72,.37,21,NAVY,align=PP_ALIGN.CENTER)
    text(s,'B',3.38,2.64,.53,.27,14,NAVY,align=PP_ALIGN.CENTER)
    text(s,'∑',3.37,2.95,.59,.59,35,NAVY,align=PP_ALIGN.CENTER)
    text(s,'i = 1',3.31,3.59,.74,.27,13,NAVY,align=PP_ALIGN.CENTER)
    equation(s,[('‖â',0,30),('i',-22000,18),(' − a',0,30),('i',-22000,18),('E',29000,18),('‖',0,30),('2',-24000,18),('2',29000,18)],4.11,2.9,3.6,.72)
    text(s,'B: batch size (up to 256) · i: a training example\naᴱ: human action · average over all 7 action values',.8,4.03,7.06,.82,16,GRAY)
    text(s,'Weight update',.78,5.13,3.1,.4,20,NAVY,True)
    equation(s,[('θ ← θ − η d',0,29),('Adam',-22000,16)],3.46,5.02,4.38,.7)
    equation(s,[('η: learning rate · d',0,15),('Adam',-22000,10),(': adaptive gradient direction',0,15)],.8,5.81,7.14,.5,size=15)
    picture(s,'loss_curves.png',8.25,2.28,4.37,2.62)
    text(s,'Actual training runs\nPale lines: training; dark lines: validation\nDotted lines: selected epochs',8.37,5.1,4.02,1.17,14,GRAY)
    line(s,.78,6.43,12.49,6.43,'D6D9DC',.8)
    text(s,'One action value: (0.15 − 0.10)² = 0.0025',.8,6.64,11.8,.37,18,TEAL)

    s = slide(6)
    text(s,'Open the notebooks in order on the prepared course server.',.78,1.48,11.8,.43,18,GRAY)
    rows=[('01','Demonstrations','Play a demo; inspect observations and actions.','20 min','Q1'),
          ('02','Policy training','Train both rates; save models, curves and evaluations.','50–60 min','Q2–Q3'),
          ('03','Evaluation','Inspect the same saved cases and position shifts.','40–50 min','Q4'),
          ('04','Failure analysis','Select ≥3 steps; mark the first error; export evidence.','20–30 min','Q5')]
    for i,(n,title,body,duration,q) in enumerate(rows):
        y=2.13+i*.92
        line(s,.78,y,12.49,y,'D6D9DC',.8)
        text(s,n,.8,y+.19,.66,.4,21,TEAL)
        text(s,title,1.74,y+.15,7.55,.36,20,NAVY,True)
        text(s,body,1.74,y+.58,7.55,.29,14,GRAY)
        text(s,duration,10.11,y+.25,1.45,.34,15,GRAY)
        text(s,q,11.71,y+.25,.79,.34,15,GRAY,align=PP_ALIGN.RIGHT)
    line(s,.78,5.81,12.49,5.81,'D6D9DC',.8)
    text(s,'Compare η = 0.001 and η = 0.0003',.8,6.04,11.7,.45,22,TEAL)
    text(s,'Keep data, initial weights, architecture, seed 17, 50 epochs, batch size 256 and device fixed.',.8,6.64,11.7,.4,14,GRAY)

    s = slide(7)
    text(s,'Two models × (5 standard + 5 shifted starts) = 20 rollouts',.78,1.47,11.8,.43,20,GRAY)
    picture(s,'positions.png',.62,2.19,5.88,4.31)
    text(s,'Stable placements: one completed example',6.96,2.24,5.56,.66,20,NAVY,True)
    for x,w,head in [(6.96,2.03,'Learning rate'),(9.46,1.36,'Standard'),(11.18,1.36,'Shifted')]:
        text(s,head,x,3.16,w,.34,15,GRAY)
    line(s,6.96,3.68,12.5,3.68,'D6D9DC',.8)
    for i,(rate,a,b) in enumerate([('0.001','2 / 5','3 / 5'),('0.0003','5 / 5','5 / 5')]):
        y=3.97+i*.79
        text(s,rate,6.96,y,2.04,.44,22)
        text(s,a,9.46,y,1.37,.44,22)
        text(s,b,11.18,y,1.37,.44,22)
    text(s,'Paired source episode and seed\nCan position shifted +0.04 m along world X',6.96,5.66,5.59,.79,16,GRAY)
    text(s,'Five cases give a coarse estimate; a position shift is not a certified out-of-distribution test.',.78,6.7,11.8,.32,14,GRAY)

    s = slide(8)
    text(s,'Locate the first visible error using frames and logged state.',.78,1.47,11.8,.43,18,GRAY)
    for i,step in enumerate([1,41,81]):
        x=.78+i*1.91
        picture(s,f'failure_step_{step:03d}.png',x,2.28,1.75,1.75)
        text(s,f'Step {step}',x,4.21,1.75,.36,16,GRAY)
    picture(s,'failure_trace.png',6.79,2.25,5.86,3.12)
    text(s,'Observed: the can remains in the source bin.',.78,5.06,5.62,.79,20)
    text(s,'In Notebook 04, select ≥3 control steps\nand mark the first observable error.',.78,6.04,5.62,.75,16,GRAY)
    text(s,'Observation → hypothesis → proposed test',6.97,5.74,5.57,.85,19,TEAL)
    text(s,'Check the offset between the requested state\nand the nearest recorded video frame.',6.97,6.44,5.53,.6,14,GRAY)

    s = slide(9)
    text(s,'Submit a 1–2 page PDF\nand the evidence ZIP',.78,2.01,5.85,1.23,28,NAVY,True)
    text(s,'Answer Q1–Q5 in your own words.\nExport the ZIP from Notebook 04.\nInclude both models and all outcomes.',.8,3.76,5.58,1.34,19)
    text(s,'The generated report is an evidence appendix.\nBlank answers or no error marker remain DRAFT.',.8,5.71,5.61,.84,15,GRAY)
    text(s,'Assessment',7.11,2.03,5.42,.43,23,NAVY,True)
    for i,(n,title) in enumerate([(20,'Demonstration understanding'),(25,'Training and experiment'),(30,'Policy evaluation'),(25,'Failure analysis and discussion')]):
        y=2.91+i*.69
        line(s,7.11,y-.12,12.49,y-.12,'D6D9DC',.8)
        text(s,title,7.11,y,4.83,.38,17)
        text(s,str(n),12.04,y,.46,.4,18,TEAL,align=PP_ALIGN.RIGHT)
    line(s,7.11,5.57,12.49,5.57,'D6D9DC',.8)
    text(s,'Success rate does not determine your grade.',7.11,6.03,5.38,.74,18)

    prs.save(HERE/'Manipulation_Learning_Lab.pptx')
    if static_output:
        s=prs.slides[1]
        movie=next(sh for sh in s.shapes if sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}videoFile') is not None)
        index=list(s.shapes._spTree).index(movie._element)
        x,y,w,h=movie.left,movie.top,movie.width,movie.height
        s.shapes._spTree.remove(movie._element)
        timing=s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}timing')
        if timing is not None:s._element.remove(timing)
        poster=s.shapes.add_picture(str(A/'success_poster.png'),x,y,w,h)
        s.shapes._spTree.remove(poster._element);s.shapes._spTree.insert(index,poster._element)
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:r.text=r.text.replace('Click to play','Video in PPTX')
        static_output=Path(static_output);static_output.parent.mkdir(parents=True,exist_ok=True);prs.save(static_output)
    notes=[f"{CONTENT['title']} — 简约版中文课堂讲解\n总时长：600 秒（10 分钟）\n"]
    for item in CONTENT['slides']:
        notes.append(f"第 {item['number']} 页 · {item['title']} · {item['seconds']} 秒\n\n{item['notes_zh']}\n")
    (HERE/'speaker_notes_zh.txt').write_text('\n\n'.join(notes))
    print('Built minimalist revision: 9 slides, 600 seconds, 3 editable equations.')


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--static-output');args=parser.parse_args()
    make_figures();make_deck(args.static_output)
