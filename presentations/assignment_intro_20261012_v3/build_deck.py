"""Build the visual classroom introduction from verified task evidence."""
import argparse
import csv
import json
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE=Path(__file__).resolve().parent
A=HERE/'assets'
C=json.loads((HERE/'slides.json').read_text())
BG='FCFBF8'; INK='252C30'; MUTED='626B70'; BLUE='3F7085'; RUST='B66C49'; LIGHT='D6DBDA'
BODY='Source Sans 3'; HEAD='Source Serif 4'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Manipulation Learning Lab'
prs.core_properties.subject='Introduction to Robotics — visual assignment introduction'
prs.core_properties.author='Manipulation Learning Lab'


def rgb(c):return RGBColor.from_string(c)


def txt(s,t,x,y,w,h,size=22,color=INK,bold=False,font=BODY,italic=False,align=None):
    b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=b.text_frame
    f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0;f.word_wrap=True
    for i,value in enumerate(t.split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph();p.text=value
        p.font.name='Source Sans 3 Semibold' if bold and font==BODY else font
        p.font.size=Pt(size);p.font.bold=False;p.font.italic=italic;p.font.color.rgb=rgb(color)
        p.space_after=Pt(4)
        if align is not None:p.alignment=align
    return b


def clean(shape):
    style=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:shape._element.remove(style)
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    return shape


def line(s,x1,y1,x2,y2,c=MUTED,width=1.4,arrow=False):
    sh=clean(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)))
    sh.line.color.rgb=rgb(c);sh.line.width=Pt(width)
    if arrow:
        end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','sm');end.set('len','sm');sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(end)
    return sh


def oval(s,x,y,w,h,c=RUST,fill=None,width=1.5):
    sh=clean(s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h)))
    sh.line.color.rgb=rgb(c);sh.line.width=Pt(width)
    if fill:sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    else:sh.fill.background()
    return sh


def box(s,x,y,w,h,c=INK,fill=None,width=1):
    sh=clean(s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)))
    sh.line.color.rgb=rgb(c);sh.line.width=Pt(width)
    if fill:sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    else:sh.fill.background()
    return sh


def photo(s,name,x,y,w,h,region=(0,0,1,1)):
    """Native crop preserves the original image bytes and aspect ratio."""
    iw,ih=Image.open(A/name).size
    left,top,right,bottom=region;rw=(right-left)*iw;rh=(bottom-top)*ih
    if rw/rh>w/h:
        keep=rh*w/h/iw;cx=(left+right)/2;left,right=cx-keep/2,cx+keep/2
    else:
        keep=rw*h/w/ih;cy=(top+bottom)/2;top,bottom=cy-keep/2,cy+keep/2
    sh=s.shapes.add_picture(str(A/name),Inches(x),Inches(y),Inches(w),Inches(h))
    sh.crop_left=left;sh.crop_right=1-right;sh.crop_top=top;sh.crop_bottom=1-bottom
    def project(px,py):return x+(px/iw-left)/(right-left)*w,y+(py/ih-top)/(bottom-top)*h
    return sh,project


def image(s,name,x,y,w,h=None):
    return s.shapes.add_picture(str(A/name),Inches(x),Inches(y),width=Inches(w),height=None if h is None else Inches(h))


def red_center(name):
    a=np.asarray(Image.open(A/name).convert('RGB')).astype(float)
    mask=(a[:,:,0]>105)&(a[:,:,0]>1.45*a[:,:,1])&(a[:,:,0]>1.5*a[:,:,2])
    yy,xx=np.where(mask);return float(np.median(xx)),float(np.median(yy))


def slide(n):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    if n!=1:txt(s,C['slides'][n-1]['title'],.7,.52,11.9,.8,34,font=HEAD)
    txt(s,str(n),12.25,7.12,.4,.24,11,MUTED,align=PP_ALIGN.RIGHT)
    item=C['slides'][n-1]
    s.notes_slide.notes_text_frame.text=f"第 {n} 页｜{item['seconds']} 秒\n{item['notes_zh']}\n\n来源：\n"+'\n'.join(item['sources'])
    return s


def charts():
    for f in (HERE/'fonts').glob('*.ttf'):font_manager.fontManager.addfont(str(f))
    plt.rcParams.update({'font.family':BODY,'font.size':15,'text.color':'#'+INK,'axes.labelcolor':'#'+MUTED,
        'xtick.color':'#'+MUTED,'ytick.color':'#'+MUTED,'axes.spines.top':False,'axes.spines.right':False,
        'axes.spines.left':False,'axes.spines.bottom':False,'figure.facecolor':'#'+BG,'axes.facecolor':'#'+BG})
    rows=list(csv.DictReader((A/'lr_a_training_log.csv').open()))
    fig,ax=plt.subplots(figsize=(4.6,3.1));epochs=[int(r['epoch']) for r in rows]
    for k,c,label in [('training_loss',BLUE,'Training'),('validation_loss',RUST,'Validation')]:
        ax.plot(epochs,[float(r[k]) for r in rows],lw=2.7,color='#'+c,label=label)
    ax.set(xlabel='Epoch',ylabel='Action error',yscale='log',xlim=(1,50));ax.legend(frameon=False,fontsize=12)
    ax.grid(axis='y',alpha=.13);fig.tight_layout();fig.savefig(A/'loss_visual.png',dpi=220);plt.close(fig)
    d=json.loads((A/'action_comparison.json').read_text())
    fig,ax=plt.subplots(figsize=(6.1,3.9));y=np.arange(3)
    for vals,dy,c,label in [(d['expert_action'],.13,BLUE,'Human'),(d['predicted_action'],-.13,RUST,'Policy')]:
        ax.barh(y+dy,vals[:3],.22,color='#'+c,label=label)
    ax.axvline(0,color='#ADB5B5',lw=1)
    ax.set(yticks=y,yticklabels=['Move X','Move Y','Move Z'],xlim=(-.2,.2),xticks=[-.2,-.1,0,.1,.2],xlabel='Normalized translation command')
    ax.invert_yaxis();ax.legend(frameon=False,ncol=2,loc='lower center',bbox_to_anchor=(.5,1.02),fontsize=17)
    fig.tight_layout();fig.savefig(A/'action_comparison.png',dpi=220);plt.close(fig)
    plt.rcParams['mathtext.fontset']='stix'
    fig=plt.figure(figsize=(10.5,1.05));fig.text(.02,.22,'$'+C['equations'][0]['latex']+'$',fontsize=37,color='#'+INK)
    fig.savefig(A/'bc_equation.png',dpi=400,transparent=True,bbox_inches='tight',pad_inches=.06);plt.close(fig)
    pts=list(csv.DictReader((A/'training_initial_positions.csv').open()))
    fig,ax=plt.subplots(figsize=(4.6,3.1))
    ax.scatter([float(p['x']) for p in pts],[float(p['y']) for p in pts],s=16,color='#A3B0B4',alpha=.7)
    d=json.loads((A/'lr_a_evaluation_summary.json').read_text())
    for condition,c in [('standard',BLUE),('shifted',RUST)]:
        xy=np.array([json.loads(e['initial_object_position'])[:2] for e in d['episodes'] if e['condition']==condition])
        ax.scatter(xy[:,0],xy[:,1],s=40,color='#'+c)
    ax.set(xlabel='Initial X (m)',ylabel='Initial Y (m)');fig.tight_layout();fig.savefig(A/'positions_visual.png',dpi=220);plt.close(fig)


def build(static_output=None):
    s=slide(1)
    txt(s,'Introduction to Robotics',.75,1.22,5.7,.4,18,MUTED)
    txt(s,'Manipulation\nLearning Lab',.72,2.0,6.0,1.75,42,font=HEAD)
    txt(s,'From Demonstrations\nto Robot Policies',.76,4.25,5.8,1.1,26,BLUE)
    photo(s,'success_step_089.png',6.95,.8,5.7,5.95)
    txt(s,'Learning Robot Manipulation from Demonstrations',.76,6.35,5.65,.54,15,MUTED)

    s=slide(2)
    txt(s,'Panda arm: pick up the can and release it in the marked bin.',.72,1.27,11.9,.3,17,MUTED)
    s.shapes.add_movie(str(A/'success_half_speed.mp4'),Inches(.7),Inches(1.57),Inches(4.9),Inches(4.9),poster_frame_image=str(A/'success_poster.png'),mime_type='video/mp4')
    txt(s,'Click to play · 15 s · half speed',.72,6.66,4.84,.35,15,MUTED)
    for i,(title,step) in enumerate([('Reach',1),('Grasp',41),('Transport',89),('Place',145)]):
        x=6.06+(i%2)*3.36;y=1.64+(i//2)*2.26
        photo(s,f'success_step_{step:03d}.png',x,y,3.07,1.64,region=(0,.19,1,.71))
        txt(s,title,x,y+1.7,3.07,.42,21)
    txt(s,'In the target. Released. Stable for 1 second.',6.08,6.43,6.44,.43,19,BLUE)
    txt(s,'Open gripper · no contact · displacement ≤1 cm',6.08,6.94,6.44,.28,15,MUTED)

    s=slide(3)
    txt(s,'A demonstration pairs the current observation with a human’s next action.',.73,1.35,11.87,.4,18,MUTED)
    s.shapes.add_movie(str(A/'human_demonstration.mp4'),Inches(.7),Inches(2.05),Inches(3.47),Inches(3.47),poster_frame_image=str(A/'human_poster.png'),mime_type='video/mp4')
    image(s,'loss_visual.png',4.7,2.45,3.78)
    photo(s,'success_step_145.png',9.13,2.05,3.47,3.47,region=(.25,.11,1,.86))
    line(s,4.29,3.69,4.61,3.69,BLUE,1.8,True);line(s,8.64,3.69,8.98,3.69,BLUE,1.8,True)
    for x,t in [(.7,'Watch a person'),(4.7,'Fit a policy'),(9.13,'Let the robot try')]:txt(s,t,x,5.75,3.55,.5,25)
    txt(s,'Demonstration · Click to play',.72,6.27,3.4,.3,14,MUTED)
    txt(s,'Training',4.7,6.27,3.55,.3,14,MUTED)
    txt(s,'Inference · 20 actions / second',9.13,6.27,3.55,.3,14,MUTED)
    txt(s,'160 train  /  20 validation  /  20 test demonstrations',4.7,6.85,7.91,.33,16,MUTED)

    s=slide(4)
    _,project=photo(s,'demonstration_state.png',.7,1.68,6.9,4.9,region=(0,.08,.95,.79))
    p=project(240,333);oval(s,p[0]-.54,p[1]-.20,1.12,.43,BLUE)
    txt(s,'Gripper pose\n+ finger positions',.96,1.93,2.6,.8,19,BLUE,italic=True)
    line(s,2.07,2.71,p[0]-.1,p[1]-.29,BLUE,1.3,True)
    p=project(*red_center('demonstration_state.png'));oval(s,p[0]-.23,p[1]-.3,.46,.61,RUST)
    txt(s,'Can pose',4.6,4.88,2.34,.48,20,RUST,italic=True);line(s,4.51,5.0,p[0]+.28,p[1],RUST,1.3,True)
    txt(s,'16',8.24,1.64,3.8,.8,54,BLUE,font=HEAD)
    txt(s,'state values',8.29,2.53,3.9,.44,24)
    txt(s,'Gripper pose (7) + finger joints (2)\n+ can pose (7)',8.29,3.07,4.2,.67,17,MUTED)
    line(s,8.3,3.88,12.29,3.88,LIGHT,1)
    txt(s,'7',8.24,4.05,3.8,.8,54,RUST,font=HEAD)
    txt(s,'commands',8.29,4.94,3.9,.44,24)
    txt(s,'3 translation + 3 rotation + 1 gripper\nNormalized commands in [−1, 1]',8.29,5.59,4.2,.75,17,MUTED)
    txt(s,'A fresh observation, 20 times per second.',.74,6.72,11.8,.36,18,MUTED)

    s=slide(5)
    txt(s,'Behavior cloning trains a neural network to predict the human’s next action.',.73,1.34,11.87,.4,18,MUTED)
    photo(s,'demonstration_state.png',.75,1.95,4.75,3.43,region=(0,.16,.67,.73))
    txt(s,'The same observation',.77,5.58,4.73,.44,22)
    txt(s,'Three hidden layers × 256 units',.77,6.05,5.3,.3,16,MUTED)
    line(s,5.6,3.69,6.14,3.69,BLUE,1.9,True)
    image(s,'action_comparison.png',6.26,1.93,6.3,4.03)
    txt(s,'Adam updates weights; validation selects the model.',6.26,6.04,6.3,.31,16,MUTED)
    image(s,'bc_equation.png',.8,6.45,6.35)
    txt(s,'Average squared error across samples\nand all 7 action values.',7.71,6.52,4.85,.58,17,MUTED)

    s=slide(6)
    for name,x,y,w,h in [('human_poster.png',.75,1.9,2.64,1.96),('failure_hd_081.png',7.19,4.47,2.64,1.96)]:
        photo(s,name,x,y,w,h,region=(0,.08,1,.85))
    image(s,'loss_visual.png',7.03,1.9,2.88,1.94)
    image(s,'positions_visual.png',.56,4.48,2.93,1.97)
    for x,y,n,t,q,description in [(3.61,2.07,'01','Watch','Q1','Inspect observations and actions.'),(10.02,2.07,'02','Train twice','Q2–Q3','Save both models and curves.'),(3.61,4.75,'03','Compare','Q4','Use the same test cases.'),(10.02,4.75,'04','Explain','Q5','Inspect at least 3 steps.')]:
        txt(s,n,x,y,2.42,.33,16,BLUE)
        txt(s,t,x,y+.51,2.55,.68,27,font=HEAD)
        txt(s,description+' '+q,x,y+1.14,2.53,.66,16,MUTED)
    txt(s,'Change only the learning rate: 0.001 or 0.0003.',.77,6.72,11.8,.34,19,BLUE)
    txt(s,'Fixed data, model, initialization, seed 17, 50 epochs, batch size 256 and device.',.77,7.08,11.3,.24,13,MUTED)

    s=slide(7)
    txt(s,'Shift only the can; pair the source episode and random seed.',.73,1.34,11.87,.4,18,MUTED)
    for i,name in enumerate(['standard_start.png','shifted_start.png']):
        x=.73+i*3.46
        _,project=photo(s,name,x,2.11,3.21,3.32,region=(0,.17,.69,.83))
        p=project(*red_center(name));oval(s,p[0]-.23,p[1]-.31,.46,.64,BLUE if i==0 else RUST)
        txt(s,'Standard' if i==0 else '+4 cm in world X',x,5.65,3.27,.5,21)
    txt(s,'Five matched starts per model',.75,6.45,6.72,.45,22,MUTED)
    txt(s,'Not a certified out-of-distribution test.',.75,6.96,6.72,.27,15,MUTED)
    txt(s,'Stable placements',8.06,1.82,4.47,.5,25,font=HEAD)
    txt(s,'Standard',8.28,2.49,1.9,.35,16,MUTED);txt(s,'Shifted',10.72,2.49,1.8,.35,16,MUTED)
    for row,(rate,n1,n2) in enumerate([('0.001',2,3),('0.0003',5,5)]):
        y=3.22+row*1.46;txt(s,'Learning rate '+rate,8.08,y-.11,4.49,.38,18)
        for x,n in [(8.22,n1),(10.65,n2)]:
            for k in range(5):oval(s,x+k*.32,y+.52,.20,.20,BLUE,BLUE if k<n else BG,1.2)
            txt(s,str(n)+'/5',x,y+.90,1.83,.34,16,MUTED)
    txt(s,'One completed example; five cases are a coarse estimate.',8.08,6.47,4.55,.57,14,MUTED)

    s=slide(8)
    txt(s,'Use frames and state logs to separate the observed error from a possible cause.',.73,1.34,11.87,.4,18,MUTED)
    for i,step in enumerate([1,41,81]):
        x=.72+i*4.2
        _,project=photo(s,f'failure_hd_{step:03d}.png',x,1.94,3.89,3.79,region=(0,.12,.64,.79))
        txt(s,f'Step {step}',x,5.92,3.88,.5,23)
        if i==2:
            p=project(*red_center(f'failure_hd_{step:03d}.png'));oval(s,p[0]-.31,p[1]-.24,.62,.48,RUST, width=2)
            txt(s,'Still here.',x+.12,2.22,2.5,.49,22,RUST,italic=True)
            line(s,x+1.28,2.75,p[0],p[1]-.3,RUST,1.4,True)
    txt(s,'Inspect ≥3 steps, mark the first observable error, then propose a test.',.74,6.67,11.82,.43,23,BLUE)

    s=slide(9)
    for x,label,sub in [(1.1,'PDF','1–2 pages · Q1–Q5'),(7.12,'ZIP','Both models · videos · logs')]:
        box(s,x,1.9,1.26,1.65,BLUE,width=1.5)
        line(s,x+.24,2.39,x+1.00,2.39,LIGHT,1.2);line(s,x+.24,2.65,x+1,2.65,LIGHT,1.2);line(s,x+.24,2.91,x+.82,2.91,LIGHT,1.2)
        txt(s,label,x+1.66,2.12,3.64,.62,36,font=HEAD)
        txt(s,sub,x+1.66,2.99,3.91,.52,22)
    txt(s,'Export the ZIP from Notebook 04. The generated report is an evidence appendix.',.8,3.83,11.8,.4,18,MUTED)
    txt(s,'Assessment',.78,4.32,11.8,.5,25,font=HEAD)
    for x,n,t in [(.8,20,'Demonstrations'),(3.92,25,'Training'),(7.04,30,'Evaluation'),(10.16,25,'Failure analysis')]:
        txt(s,str(n),x,5.04,2.5,.71,39,BLUE,font=HEAD);txt(s,t,x,5.91,2.6,.45,21)
    txt(s,'A well-explained failure can earn full credit.',.8,6.74,11.8,.41,23)

    prs.save(HERE/'Manipulation_Learning_Lab.pptx')
    if static_output:
        for s in prs.slides:
            for sh in list(s.shapes):
                if sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}videoFile') is not None:
                    name='success_poster.png' if s==prs.slides[1] else 'human_poster.png'
                    idx=list(s.shapes._spTree).index(sh._element);x,y,w,h=sh.left,sh.top,sh.width,sh.height;s.shapes._spTree.remove(sh._element)
                    pic=s.shapes.add_picture(str(A/name),x,y,w,h);s.shapes._spTree.remove(pic._element);s.shapes._spTree.insert(idx,pic._element)
            timing=s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}timing')
            if timing is not None:s._element.remove(timing)
            for sh in s.shapes:
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:r.text=r.text.replace('Click to play','Video in PPTX')
        path=Path(static_output);path.parent.mkdir(parents=True,exist_ok=True);prs.save(path)
    notes=['Manipulation Learning Lab — 以画面讲解的 v3\n9 页，600 秒。\n']
    for item in C['slides']:notes.append(f"第 {item['number']} 页 · {item['title']} · {item['seconds']} 秒\n\n{item['notes_zh']}\n")
    (HERE/'speaker_notes_zh.txt').write_text('\n\n'.join(notes))
    print('Built 9 visual slides with two embedded real videos and one typeset equation.')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--static-output');args=p.parse_args();charts();build(args.static_output)
