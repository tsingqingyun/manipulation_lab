"""Delivery acceptance for the visual revision; no learning experiments rerun."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import subprocess
import zipfile
import imageio.v2 as imageio
import imageio_ffmpeg
import numpy as np
import pymupdf
from PIL import Image
from pptx import Presentation

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[1];A=HERE/'assets'


def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def words(text):return re.findall(r"[A-Za-z0-9]+(?:[’'-][A-Za-z0-9]+)*",text)


def main(browser):
    content=json.loads((HERE/'slides.json').read_text())
    deck=Presentation(HERE/'Manipulation_Learning_Lab.pptx')
    assert len(deck.slides)==len(content['slides'])==9
    assert sum(x['seconds'] for x in content['slides'])==600
    assert len(content['equations'])==1
    counts=[];texts=[]
    for i,s in enumerate(deck.slides):
        t=' '.join(sh.text for sh in s.shapes if sh.has_text_frame);texts.append(t)
        assert not re.search('[\u4e00-\u9fff]',t)
        assert content['slides'][i]['notes_zh'] in s.notes_slide.notes_text_frame.text
        for sh in s.shapes:
            assert sh.left>=0 and sh.top>=0,(i,sh.name)
            assert sh.left+sh.width<=deck.slide_width+10 and sh.top+sh.height<=deck.slide_height+10,(i,sh.name)
        counts.append(len(words(t)))
    previous=Presentation(HERE.parent/'assignment_intro_20261012_v2/Manipulation_Learning_Lab.pptx')
    old_counts=[len(words(' '.join(sh.text for sh in s.shapes if sh.has_text_frame))) for s in previous.slides]
    # The user's live revision requested explanations back: retain a reduction,
    # but do not enforce the earlier, overly terse draft's 35% reduction target.
    assert sum(counts)<sum(old_counts)
    (HERE/'visible_slide_text.txt').write_text('\n\n'.join(f'Slide {i+1}\n{t}' for i,t in enumerate(texts))+'\n')
    media_names=['success_half_speed.mp4','human_demonstration.mp4']
    with zipfile.ZipFile(HERE/'Manipulation_Learning_Lab.pptx') as z:
        media=[n for n in z.namelist() if n.endswith('.mp4')]
        assert len(media)==2
        assert {hashlib.sha256(z.read(n)).hexdigest() for n in media}=={sha(A/n) for n in media_names}
        formula=[n for n in z.namelist() if n.startswith('ppt/media/') and hashlib.sha256(z.read(n)).hexdigest()==sha(A/'bc_equation.png')]
        assert len(formula)==1
    visual=json.loads((HERE/'visual_provenance.json').read_text())
    for f in visual['sources']:assert sha(ROOT/f['path'])==f['sha256']
    for name,h in visual['new_asset_sha256'].items():assert sha(A/name)==h
    assert max(visual['failure_replay']['max_absolute_state_errors'].values())<1e-6
    for x in visual['paired_starts']:assert x['max_state_error']<1e-6
    d=json.loads((A/'action_comparison.json').read_text())
    mse=float(np.mean((np.asarray(d['predicted_action'])-np.asarray(d['expert_action']))**2))
    assert abs(mse-d['mean_squared_error'])<1e-10
    assert len(d['expert_action'])==len(d['predicted_action'])==7
    equation=Image.open(A/'bc_equation.png')
    assert equation.width>2000 and equation.height>200
    videos=[]
    for name in media_names+['failure_original.mp4']:
        subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(),'-v','error','-i',str(A/name),'-f','null','-'],check=True,capture_output=True)
        r=imageio.get_reader(A/name);m=r.get_meta_data();n=sum(1 for _ in r);r.close()
        videos.append({'file':name,'seconds':m['duration'],'size':m['size'],'decoded_frames':n})
    pdf=pymupdf.open(HERE/'Manipulation_Learning_Lab.pdf');assert len(pdf)==9
    fonts={f[3].split('+')[-1] for page in pdf for f in page.get_fonts()}
    assert any('SourceSans3' in f for f in fonts) and any('SourceSerif4' in f for f in fonts)
    (HERE/'preview').mkdir(exist_ok=True)
    sheet=Image.new('RGB',(1920,1080),'#FCFBF8');pdf_counts=[]
    for i,page in enumerate(pdf):
        t=page.get_text();assert not re.search('[\u4e00-\u9fff]',t)
        assert not list(page.annots() or [])
        expected=texts[i].replace('Click to play','Video in PPTX')
        missing=set(words(expected.lower()))-set(words(t.lower()));assert not missing,(i+1,missing)
        pdf_counts.append(len(words(t)))
        dest=HERE/f'preview/slide_{i+1:02d}.png';page.get_pixmap(matrix=pymupdf.Matrix(1.4,1.4)).save(dest)
        im=Image.open(dest).convert('RGB');im.thumbnail((640,360));sheet.paste(im,((i%3)*640,(i//3)*360))
    sheet.save(HERE/'preview/contact_sheet.png')
    playback={}
    if browser:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b=p.chromium.launch(executable_path=browser,headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
            page=b.new_page()
            for name in media_names:
                page.goto((A/name).as_uri());page.wait_for_function('document.querySelector("video")?.readyState >= 2')
                page.evaluate('document.querySelector("video").muted=true')
                page.evaluate('document.querySelector("video").currentTime=0')
                page.evaluate('document.querySelector("video").play()');page.wait_for_function('document.querySelector("video").currentTime>.5')
                page.evaluate('document.querySelector("video").currentTime=document.querySelector("video").duration-.5')
                page.wait_for_function('document.querySelector("video").ended')
                playback[name]=page.evaluate('(()=>{const v=document.querySelector("video");return {duration:v.duration,ended:v.ended,error:v.error};})()')
                assert playback[name]['ended'] and playback[name]['error'] is None
            playback['browser_version']=b.version;b.close()
    contract=json.loads((HERE.parent/'assignment_intro_20261012_v2/source_contract.json').read_text())
    contract['files']={p:sha(ROOT/p) for p in contract['files']}
    contract['base_commit']=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
    contract['revision']='Visual v3; supplementary same-observation example, not a new policy benchmark'
    (HERE/'source_contract.json').write_text(json.dumps(contract,indent=2)+'\n')
    result={'passed':True,'slides':9,'notes_seconds':600,'visible_language':'English','notes_language':'Chinese',
            'native_text_word_counts':counts,'v2_native_text_word_counts':old_counts,
            'native_text_reduction_fraction':1-sum(counts)/sum(old_counts),
            'count_scope':'Editable slide text; mathematical graphic and chart-internal labels are excluded from both counts.',
            'pdf_text_coverage':True,'fonts':sorted(fonts),'media':videos,'embedded_media_matches':True,'playback':playback,
            'failure_replay_errors':visual['failure_replay']['max_absolute_state_errors'],
            'single_state_action_mse':mse,'formula':'Mean of componentwise squared differences across all seven components and batch samples.',
            'formula_rendering':'STIX high-resolution image with LaTeX source; not individual manually placed glyphs.',
            'limitations':['Native PowerPoint playback on the instructor computer remains untested.',
                           'Same-state illustration is not a new held-out rollout or policy performance estimate.',
                           'PDF is static. Equation is a typeset image; LaTeX source is supplied.']}
    result['deliverable_sha256']={n:sha(HERE/n) for n in ['Manipulation_Learning_Lab.pptx','Manipulation_Learning_Lab.pdf','speaker_notes_zh.txt']}
    (HERE/'verification.json').write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps({'passed':True,'slides':9,'native_words':sum(counts),'previous_words':sum(old_counts),'videos':len(media_names),'browser_playback':bool(playback)},indent=2))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--browser');main(p.parse_args().browser)
