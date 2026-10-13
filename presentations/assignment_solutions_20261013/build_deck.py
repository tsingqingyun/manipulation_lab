"""Build an editable solution deck from this run's recorded evidence, without simulation."""
import argparse
import csv
import hashlib
import json
import shutil
from pathlib import Path

import imageio.v2 as imageio
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
A = HERE / "assets"
BG, INK, MUTED, BLUE, RUST, LIGHT = "FCFBF8", "252C30", "626B70", "3F7085", "B66C49", "E9ECE9"
BODY, HEAD = "Source Sans 3", "Source Serif 4"


def read(path):
    return json.loads(Path(path).read_text())


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write(path, value):
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n")


def txt(s, text, x, y, w, h, size=22, color=INK, bold=False, font=BODY):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = b.text_frame
    f.word_wrap = True
    f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph()
        p.text = line
        p.font.name = "Source Sans 3 Semibold" if bold and font == BODY else font
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(8)
    return b


def rect(s, x, y, w, h, fill=LIGHT):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill); b.line.fill.background()
    return b


def picture(s, name, x, y, w, h):
    p = A / name
    iw, ih = Image.open(p).size
    scale = min(w / iw, h / ih)
    return s.shapes.add_picture(str(p), Inches(x + (w-iw*scale)/2),
                               Inches(y + (h-ih*scale)/2), width=Inches(iw*scale), height=Inches(ih*scale))


def table(s, headers, rows, widths, x=.72, y=1.85, row_h=.66, size=20):
    for i, row in enumerate([headers] + rows):
        xx = x
        for val, width in zip(row, widths):
            rect(s, xx, y+i*row_h, width-.025, row_h-.025, BLUE if i == 0 else (LIGHT if i % 2 else "F3F3EF"))
            txt(s, str(val), xx+.12, y+i*row_h+.12, width-.22, row_h-.15,
                size, "FFFFFF" if i == 0 else INK, i == 0)
            xx += width


def prepare(run):
    experiment = Path(read(run / "experiment_current.json")["path"])
    analysis = Path(read(run / "analysis_current.json")["path"])
    demo = Path(read(run / "demonstration_current.json")["path"])
    e = read(experiment / "experiment.json")
    assert e["complete"] and all(e["checks"].values())
    data = {"run": str(run), "experiment": e, "variants": {}, "failure": read(analysis / "failure_analysis.json"),
            "facts": read(analysis / "observed_facts.json"), "answers": read(HERE / "answers.json")}
    sources = []
    def copy(src, dst):
        shutil.copyfile(src, dst)
        sources.append({"source": str(src), "file": str(dst.relative_to(HERE)), "sha256": sha(dst)})
    for v, p in e["variants"].items():
        tr, ev = experiment / p["training"], experiment / p["evaluation"]
        summary = read(ev / "summary.json")
        assert summary["errors"] == 0 and len(summary["episodes"]) == 10
        data["variants"][v] = {"train": read(tr / "training_summary.json"), "eval": summary,
                               "history": list(csv.DictReader((tr / "training_log.csv").open()))}
        for src, name in [(tr / "training_log.csv", "training_log.csv"), (tr / "training_summary.json", "training_summary.json"),
                          (tr / "training_config.json", "training_config.json"), (ev / "summary.json", "evaluation_summary.json"),
                          (ev / "paired_results.csv", "paired_results.csv"), (ev / "test_manifest.json", "test_manifest.json")]:
            copy(src, HERE / "evidence" / f"{v}_{name}")
    primary = experiment / e["variants"]["lr_a"]["evaluation"]
    copy(experiment / "experiment.json", HERE / "evidence" / "experiment.json")
    copy(primary / "evaluation_config.json", HERE / "evidence" / "evaluation_config.json")
    copy(demo / "dataset_inspection.json", HERE / "evidence" / "dataset_inspection.json")
    copy(analysis / "observed_facts.json", HERE / "evidence" / "observed_facts.json")
    copy(analysis / "failure_analysis.json", HERE / "evidence" / "failure_analysis.json")
    data["inspection"] = read(demo / "dataset_inspection.json")
    failure = next(r for r in data["variants"]["lr_a"]["eval"]["episodes"] if r["episode_id"] == data["facts"]["episode"])
    success = next(r for r in data["variants"]["lr_a"]["eval"]["episodes"] if r["stable_success"])
    data["success_episode"] = success; data["failure_episode"] = failure
    for label, row in [("failure", failure), ("success", success)]:
        copy(primary / row["video"], A / f"{label}.mp4")
        copy(primary / f'{row["episode_id"]}.npz', HERE / "evidence" / f"{label}_trace.npz")
        with imageio.get_reader(A / f"{label}.mp4") as reader:
            count = reader.count_frames()
            imageio.imwrite(A / f"{label}_poster.png", reader.get_data(min(count-1, 12)))
            if label == "success":
                imageio.imwrite(A / "success_place.png", reader.get_data(max(0, count-4)))
    copy(demo / "demonstration.mp4", A / "demonstration.mp4")
    with imageio.get_reader(A / "demonstration.mp4") as reader:
        n = reader.count_frames()
        for i, index in enumerate([0, n//2, n-1]):
            imageio.imwrite(A / f"demo_{i}.png", reader.get_data(index))
    for frame in data["facts"]["keyframes"]:
        copy(analysis / frame["file"], A / frame["file"])
    write(HERE / "media_provenance.json", {"run": str(run), "files": sources,
          "derived_frames": "Unaltered decoded frames from the copied MP4s; no generated robot images.",
          "claims": "This single fixed-seed run; no real-robot or independent multi-seed claim."})
    write(HERE / "results.json", data)
    return data


def charts(d):
    for f in (HERE / "fonts").glob("*.ttf"):
        font_manager.fontManager.addfont(str(f))
    plt.rcParams.update({"font.family": BODY, "font.size": 14, "figure.facecolor": "#"+BG,
        "axes.facecolor": "#"+BG, "axes.spines.top": False, "axes.spines.right": False,
        "axes.labelcolor": "#"+INK, "text.color": "#"+INK})
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.25), sharey=True)
    for ax, (v, entry) in zip(axes, d["variants"].items()):
        h = entry["history"]; ep = [int(r["epoch"]) for r in h]
        for field, label, color in [("training_loss", "Training", BLUE), ("validation_loss", "Validation", RUST)]:
            ax.plot(ep, [float(r[field]) for r in h], label=label, color="#"+color, lw=2.4)
        best = entry["train"]["best_epoch"]
        ax.axvline(best, color="#"+MUTED, ls="--", label=f"Selected epoch {best}")
        ax.set(xlabel="Epoch", ylabel="Mean squared action error", yscale="log", xlim=(1,50),
               title=f'{v}  |  learning rate {d["experiment"]["variants"][v]["learning_rate"]}')
        ax.legend(fontsize=11); ax.grid(alpha=.15)
    fig.tight_layout(); fig.savefig(A / "losses.png", dpi=200); plt.close(fig)
    inspection = d["inspection"]
    lookup = {e["episode"]: e["initial_object_position_m"] for e in inspection["episodes"]}
    tr = np.array([lookup[e] for e in inspection["splits"]["train"]])
    fig, ax = plt.subplots(figsize=(6.5,4.8))
    ax.scatter(tr[:,0], tr[:,1], c="#A3B0B4", s=23, alpha=.65, label="160 training starts")
    sets = []
    for condition, color, marker in [("standard", BLUE, "o"), ("shifted", RUST, "x")]:
        xy=np.array([json.loads(r["initial_object_position"]) for r in d["variants"]["lr_a"]["eval"]["episodes"] if r["condition"]==condition])
        ax.scatter(xy[:,0],xy[:,1],s=65,c="#"+color,marker=marker,label=condition.capitalize());sets.append(xy)
    for p,q in zip(*sets):ax.plot([p[0],q[0]],[p[1],q[1]],color="#"+MUTED,lw=.8)
    ax.set(xlabel="Initial can world X (m)",ylabel="Initial can world Y (m)");ax.legend(fontsize=11)
    ax.set_aspect("equal",adjustable="datalim");fig.tight_layout();fig.savefig(A / "positions.png",dpi=200);plt.close(fig)
    trace=np.load(HERE / "evidence" / "failure_trace.npz")
    fig,axes=plt.subplots(4,1,figsize=(7.5,5.6),sharex=True)
    t=np.arange(101)
    axes[0].plot(t,trace["eef_position"][:101,2],color="#"+BLUE,label="End effector Z")
    axes[0].plot(t,trace["object_position"][:101,2],color="#"+RUST,label="Can Z")
    axes[0].set_ylabel("World Z (m)");axes[0].legend(fontsize=10,loc="upper left",ncol=2)
    axes[1].plot(t,trace["gripper_width"][:101],color="#"+BLUE);axes[1].set_ylabel("Opening (m)")
    axes[2].step(np.arange(1,101),trace["actions"][:100,-1],where="pre",color="#"+RUST);axes[2].set_ylabel("Close command")
    axes[3].step(t,trace["gripper_contact"][:101].astype(int),where="post",color="#"+BLUE);axes[3].set(ylabel="Contact",xlabel="Control step (20 Hz)",yticks=[0,1])
    for ax in axes:
        ax.axvline(d["answers"]["FIRST_ERROR_STEP"],color="crimson",ls="--",lw=1.5);ax.set_xlim(0,100);ax.grid(alpha=.12)
    fig.tight_layout();fig.savefig(A / "failure_zoom.png",dpi=200);plt.close(fig)


def build(d, static=False):
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    prs.core_properties.title="Manipulation Learning Lab — Worked Solutions"
    prs.core_properties.author="Manipulation Learning Lab"
    records=[]
    def slide(title, tag, notes, sources, seconds=60):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
        n=len(prs.slides)
        txt(s,tag.upper(),.72,.38,11.8,.28,12,BLUE,bold=True)
        txt(s,title,.72,.91,11.9,.88,32,font=HEAD)
        txt(s,"MANIPULATION LEARNING LAB  /  WORKED SOLUTIONS  /  13 OCT 2026",.72,7.1,11.8,.25,10,MUTED)
        txt(s,f"{n:02d}",12.15,7.06,.4,.3,12,MUTED)
        s.notes_slide.notes_text_frame.text=f"第 {n} 页｜建议 {seconds} 秒\n{notes}\n\n证据来源：\n"+"\n".join(sources)
        records.append({"slide":n,"title":title,"seconds":seconds,"notes_zh":notes,"sources":sources})
        return s
    def takeaway(s,text):
        rect(s,.72,6.34,11.88,.53)
        txt(s,text,.88,6.46,11.52,.32,17,BLUE,bold=True)
    def movie(s,label,x,y,w,h):
        if static:picture(s,label+"_poster.png",x,y,w,h)
        else:s.shapes.add_movie(str(A/(label+".mp4")),Inches(x),Inches(y),Inches(w),Inches(h),poster_frame_image=str(A/(label+"_poster.png")),mime_type="video/mp4")
    a,b=d["variants"]["lr_a"],d["variants"]["lr_b"]
    s=slide("From demonstration to a defensible answer","Instructor solution", "这份题解沿着四份 Notebook 的数据流讲解五道问题。所有数字来自本次真实执行：先检查示教数据，再训练两种学习率，最后进行二十次闭环仿真。请强调这是一份有完整证据的参考解答，不是要求每位学生取得相同成功率。学生要学会解释输入输出、控制变量、结果与原因之间的界限。",["evidence/experiment.json","results.json"],35)
    txt(s,"Manipulation\nLearning Lab",.75,2.05,6.25,1.9,43,font=HEAD)
    txt(s,"Worked solutions to Q1–Q5",.78,4.13,6.2,.55,25,BLUE)
    txt(s,"4 executed notebooks\n2 policies × 50 epochs\n20 real simulator rollouts",.78,5.02,6.1,1.13,22)
    picture(s,"success_place.png",7.5,1.77,4.6,4.6)
    s=slide("What actually counts as task success?","Task and evidence", "先播放本次主要策略 lr_a 的真实成功视频。机器人需要抓取、搬运、松开并稳定放置，不能把抬起来等同于任务完成。环境成功是执行期间曾满足环境标志；稳定放置还要求目标区域、开口至少五点五厘米、无夹爪接触，继续执行策略一秒，位移不超过一厘米、速度不超过每秒五厘米。这里是持续运行同一个策略，不是人为把物体静止下来。",["assets/success.mp4","evidence/evaluation_config.json","src/manipulation_lab/evaluation.py"],60)
    movie(s,"success",.75,1.83,4.15,4.15)
    txt(s,f'Primary policy lr_a · {d["success_episode"]["episode_id"]}\nClick to play the recorded rollout',.78,6.02,4.2,.5,15,MUTED)
    txt(s,"1  Reach and grasp\n2  Transport to the target\n3  Open the gripper; remove contact\n4  Remain stable for 1 second",5.5,2.0,6.75,2.4,25)
    txt(s,"Stable release: opening ≥ 0.055 m\nDisplacement ≤ 0.01 m; speed ≤ 0.05 m/s\nThe policy keeps running during the check.",5.5,4.76,6.7,1.15,20,BLUE)
    s=slide("Q1 · The policy observes 16 state values","Notebook 01 / observations", "观测是决策时已知的状态，动作是下一步给控制器的命令。逐项相加：末端位置三维、四元数四维、夹爪两个关节位置两维、罐子位置三维、四元数四维，共十六维。位置使用世界坐标系和米，Z 轴向上；四元数顺序是 xyzw。原始物体向量十四维不等于策略使用十四维；当前适配器仅取 object[7:14] 的世界位姿，训练和执行时一致。夹爪两个关节位置是两维观测，不意味着两个独立策略动作。",["notebooks/01_demonstrations.ipynb","src/manipulation_lab/common.py","evidence/dataset_inspection.json"])
    table(s,["Observation","Size","Meaning / units"],[["End-effector position","3","World XYZ, meters"],["End-effector orientation","4","Quaternion, xyzw"],["Finger joint positions","2","Two measured slide joints, meters"],["Can position + orientation","7","World XYZ + xyzw quaternion"]],[4.25,1,6.63],row_h=.75)
    takeaway(s,"3 + 4 + 2 + 7 = 16. The adapter uses object[7:14] in both training and rollout.")
    s=slide("Q1 · Seven commands do not mean seven positions","Notebook 01 / actions", "动作七维分为三个平移增量、三个轴角旋转增量和一个夹爪命令。每维是归一化命令，不可直接把零点一解释成零点一米。平移缩放上限为五厘米，旋转为零点五弧度；例如平移命令零点二对应控制器请求的一厘米增量，实际末端运动仍受控制器和物理约束影响。夹爪命令正值闭合、负值张开；两个手指关节由同一个开合动作驱动。二十赫兹意味着每步五十毫秒，步号不是采集时间戳。",["configs/environment.yaml","notebooks/01_demonstrations.ipynb","src/manipulation_lab/simulator.py"])
    table(s,["Action channels","Normalized command","Controller interpretation"],[["Translation  (3)","[-1, 1]","World increments; limits ±0.05 m"],["Rotation  (3)","[-1, 1]","Axis-angle increments; limits ±0.5 rad"],["Gripper  (1)","Positive / negative","Close / open; unitless"]],[3.45,3.35,5.08],row_h=.78)
    txt(s,"Example: translation command 0.2 requests a 0.01 m controller increment.\nThe measured displacement can differ because of dynamics and contact.",.85,5.18,11.55,.9,20)
    takeaway(s,"Finger joints = 2   /   Gripper observation values = 2   /   Gripper action = 1   /   20 Hz")
    s=slide("Q1 · A trajectory explains what the final image cannot","Notebook 01 / demonstration", "这里展示同一条真实人类示教 demo_0 的开始、中间和结束画面。最终照片只能说明某个时刻的状态，不能说明何时闭合、抬升之前是否抓稳、何时释放，或者动作和下一时刻状态的因果顺序。完整轨迹提供观测与专家动作的配对，BC 才能学习状态到动作的映射。示教回放恢复记录的模拟器状态，不是当前学习策略运行成功的证据。",["assets/demonstration.mp4","notebooks/01_demonstrations.ipynb"],45)
    for i,label in enumerate(["Recorded start","Middle of demonstration","Recorded end"]):
        picture(s,f"demo_{i}.png",.75+i*4.03,1.95,3.8,3.8)
        txt(s,label,.8+i*4.03,5.71,3.8,.4,19,BLUE)
    takeaway(s,"Sequence reveals approach, closure, lift and release. Demonstration playback is recorded evidence.")
    s=slide("Q2 · Behavior cloning learns the expert action","Notebook 02 / learning rule", "行为克隆就是监督学习：输入当前十六维状态，输出七维动作，以对应的专家动作为标签，最小化所有动作分量的均方误差。优化器根据梯度更新网络权重。每个 epoch 遍历完整训练样本。归一化均值和标准差只用训练集计算；验证集只用于比较每轮模型，选验证误差最小的检查点；测试集不训练也不选模型。整条轨迹属于同一个划分，避免相邻高度相关的状态跨集合泄漏。",["src/manipulation_lab/training.py","src/manipulation_lab/policy.py","evidence/dataset_inspection.json"])
    for x,w,title,body in [(.75,3.4,"Observation","16 measured state values\nTraining-only normalization"),(4.9,3.45,"BC policy","MLP: 256 → 256 → 256\nPredict 7 action values"),(9.1,3.45,"Expert target","Recorded action a*\nCompare with prediction")]:
        rect(s,x,2.0,w,1.65);txt(s,title,x+.17,2.16,w-.3,.4,24,BLUE,bold=True);txt(s,body,x+.17,2.76,w-.3,.7,18)
    txt(s,"→",4.31,2.56,.5,.5,32,BLUE);txt(s,"→",8.52,2.56,.5,.5,32,BLUE)
    txt(s,"Loss = mean of (predicted action − expert action)²",.85,4.12,11.5,.65,28,font=HEAD)
    txt(s,"200 full episodes: 160 train / 20 validation / 20 test\n23,207 recorded control steps; test rollouts never update weights.",.85,5.14,11.5,.87,23)
    takeaway(s,"Gradient update learns weights; validation loss selects the checkpoint; test cases measure behavior.")
    s=slide("Q2 · Change one variable; keep the comparison fair","Notebook 02 / controlled experiment", "两次训练唯一计划改变的变量是学习率。数据、整轨迹划分、网络、初始化、随机种子十七、五十轮预算、批大小二百五十六和设备均保持一致。代码还比较初始权重哈希、数据与评测清单，所有检查必须通过。本次使用 CPU 训练、软件渲染；最初 EGL 失败发生在示教渲染阶段，未进入训练，日志保留。lr_a 在查看测试结果前就固定为主要分析模型，所以即使另一组分数更高，也不能事后换掉主要模型。",["evidence/experiment.json","evidence/lr_a_training_config.json","evidence/lr_b_training_config.json"])
    table(s,["Setting","lr_a","lr_b"],[["Learning rate","0.001","0.0003"],["Seed / epochs / batch","17 / 50 / 256","17 / 50 / 256"],["Hidden layers","256 / 256 / 256","256 / 256 / 256"],["Selected validation epoch",str(a["train"]["best_epoch"]),str(b["train"]["best_epoch"])],["Test cases","Same 5 standard + 5 shifted","Same 5 standard + 5 shifted"]],[4.05,3.915,3.915],row_h=.65,size=19)
    takeaway(s,"All control checks passed. lr_a was fixed as the primary policy before test evaluation.")
    s=slide("Q3 · Read training and validation curves together","Notebook 02 / convergence", "先解释横轴是训练轮次，纵轴是对归一化动作计算的均方误差，使用对数刻度。蓝色训练、棕色验证，虚线标出最低验证误差的轮次。不要看到第五十轮训练误差最低就选择第五十轮：本作业按验证集选择权重。训练误差持续下降而验证误差没有同步下降，说明在训练示例上拟合更好，不等于新状态上的表现改善。两组使用相同纵轴，便于直接比较；所有数值保留在原始 CSV。",["evidence/lr_a_training_log.csv","evidence/lr_b_training_log.csv"])
    picture(s,"losses.png",.75,1.8,11.8,4.24)
    takeaway(s,f'Best validation MSE: lr_a = {a["train"]["best_validation_mse"]:.5f}; lr_b = {b["train"]["best_validation_mse"]:.5f}. Both trained for 50 epochs.')
    s=slide("Q3 · Lower training error did not yield more successes","Notebook 03 / observed results", "直接比较同一次运行的真实数字。lr_a 第五十轮训练误差更低，但用于评测的是各自验证集选择的检查点；其标准条件和偏移条件均只有三次成功，lr_b 均为五次。由此不能用训练误差代替闭环任务指标。注意这不是在同一权重上只改变误差的因果实验，也不是统计意义上证明低学习率普遍更优。每种条件只有五个案例，一个案例就是二十个百分点；只有一个训练种子，配对的标准和偏移也不是独立样本。",["evidence/lr_a_training_log.csv","evidence/lr_b_training_log.csv","evidence/lr_a_evaluation_summary.json","evidence/lr_b_evaluation_summary.json"])
    rows=[]
    for v,entry in d["variants"].items():
        cond=entry["eval"]["conditions"]
        rows.append([v,f'{float(entry["history"][-1]["training_loss"]):.5f}',f'{entry["train"]["best_validation_mse"]:.5f}',f'{round(cond["standard"]["stable_success_rate"]*5)}/5',f'{round(cond["shifted"]["stable_success_rate"]*5)}/5'])
    table(s,["Variant","Epoch-50 train","Selected val.","Standard","Shifted"],rows,[1.5,2.95,2.85,2.29,2.29],row_h=.85,size=22)
    txt(s,"A rollout changes the states the policy will see next.\nSmall action errors can accumulate; a missed grasp changes the rest of the task.",.85,4.8,11.5,1.12,24)
    takeaway(s,"Counts use stable released placement. Environment-success counts happen to match in this run.")
    s=slide("Q4 · Position Shift changes only the initial can position","Notebook 03 / coverage", "散点图每一点是一条示教或一次测试的初始位置，不是运动轨迹；颜色也不表示成功失败。灰点为一百六十条训练轨迹的起点，蓝点是测试的标准起点，叉号是沿世界 X 正方向平移四厘米的起点，细线连接同一个配对案例。配对使用相同的源轨迹和随机种子，偏移后仍须在源箱内。这只是受控位置扰动，二维图无法证明完整十六维输入分布之外，因此不能称严格的分布外测试。",["evidence/lr_a_test_manifest.json","evidence/dataset_inspection.json","src/manipulation_lab/teaching.py"])
    picture(s,"positions.png",.7,1.75,7.0,4.57)
    txt(s,"Standard start → +0.04 m in world X\nSame source state and evaluation seed\nShifted can stays inside the source bin",8.03,2.2,4.55,1.8,22)
    txt(s,"Dots are initial positions.\nLines connect paired starts.\nNeither marks a trajectory or success.",8.03,4.45,4.55,1.4,21,BLUE)
    takeaway(s,"2D overlap and a four-centimeter shift do not certify out-of-distribution generalization.")
    s=slide("Q4 · Same success labels, different completion times","Notebook 03 / paired cases", "本次所有配对案例的成功失败标签都没有因偏移而改变，但完成时间发生变化，不能把总成功率相同解释为行为完全不变。表内显示结果和执行步数。主要策略在 demo_185、种子一百零二上两种条件均成功，但标准需要一百四十一步，偏移后一百五十六步，增加十五步即零点七五秒。lr_b 在 demo_48 上反而从一百八十二步减少到一百四十一步。这些是观察事实；位置影响抓取几何或状态区域是可能机制，表格本身不能诊断原因。五百步表示未在时间预算内完成，而不是成功时间。",["evidence/lr_a_paired_results.csv","evidence/lr_b_paired_results.csv"])
    paired={}
    for v,en in d["variants"].items():
        for r in en["eval"]["episodes"]:
            key=(r["source_episode"],r["seed"]);paired.setdefault(key,{})[v+"_"+r["condition"]]=("Pass" if r["stable_success"] else "Fail")+f' · {r["steps"]}'
    table(s,["Source / seed","lr_a std.","lr_a shift","lr_b std.","lr_b shift"],[[f"{k[0]} / {k[1]}"]+[p[v+"_"+c] for v in ["lr_a","lr_b"] for c in ["standard","shifted"]] for k,p in paired.items()],[3.2,2.17,2.17,2.17,2.17],row_h=.64,size=19)
    takeaway(s,"Cells show result · control steps. lr_a / demo_185: 141 → 156 steps (+0.75 s); both succeed.")
    s=slide("Q5 · The arm moves on without the can","Notebook 04 / real failure", "播放主要策略第一条失败 standard_00。最明显的结果是机械臂继续搬运动作，但罐子留在源箱。不要把第五百步最终没放进去标为第一错误，那只是结果。我们沿状态和接触记录向前检查：四十九至五十六步有接触，五十七步首次失去接触，同时闭合命令接近正一；后续一直没有重新接触。把五十七步标为最早有明确数据支持的抓持丢失，不把它冒充最早的因果起点。更早的对齐偏差是否已经构成错误，还需要对照。",["assets/failure.mp4","evidence/failure_trace.npz","evidence/observed_facts.json"],70)
    movie(s,"failure",.75,1.85,4.2,4.2)
    txt(s,"lr_a · standard_00 · 500 steps / 25 s\nClick to play the unmodified failure video",.78,6.05,4.3,.5,15,MUTED)
    txt(s,"Contact is present at steps 49–56.\nAt step 57, contact becomes false.\nThe gripper still receives a close command.\nThe arm rises; the can stays behind.",5.55,2.07,6.68,2.85,24)
    txt(s,"Failure stage: grasp retention / start of lift\nAnnotation: first supported loss of contact\nRoot cause: still a hypothesis",5.55,5.1,6.68,1.0,20,BLUE)
    s=slide("Q5 · Pin the explanation to recorded steps","Notebook 04 / keyframes", "四张图分别是第四十五、五十三、五十七和六十五步。它们都恰好有实际视频帧，偏移为零，帧号分别十一、十三、十四和十六。请区分动作和状态：a[t-1] 导致 s[t]；视频每四步在执行动作之后记录，所以第一帧是状态一而不是状态零。第五十七步为二点八五秒控制经过时间，不是采集时间戳。普通情况下请求步号未必恰好有视频帧，必须报告偏移，不能把邻帧图像当成该时刻精确状态。",["evidence/observed_facts.json","evidence/failure_trace.npz"])
    for i,f in enumerate(d["facts"]["keyframes"]):
        x=.74+i*3.03
        picture(s,f["file"],x,1.85,2.9,2.65)
        txt(s,f'Step {f["requested_control_step"]}  /  frame {f["video_frame_index"]}',x+.03,4.59,2.85,.38,20,BLUE,bold=True)
        st=f["requested_state"]
        txt(s,f'Opening {st["gripper_opening_m"]:.4f} m\nCan Z {st["object_world_position_m"][2]:.4f} m\nFrame offset {f["offset_steps"]:+d} steps',x+.03,5.08,2.85,1.05,17)
    takeaway(s,"a[t−1] leads to s[t]. Frames are sampled after actions: steps 1, 5, 9, …; use the saved mapping.")
    s=slide("Q5 · Separate measured facts from a causal story","Notebook 04 / diagnosis", "左图放大前一百步，分别画世界高度、夹爪开口、归一化闭合命令和接触布尔值，避免混用单位。红线是人工标注，不是算法自动识别的根因。第五十七步开口约二点二四厘米，之后夹爪几乎完全闭合，末端明显升高，而罐子最高只比初始高约五点七毫米。事实支持没有形成可持续的抓持。可能是进入位置或姿态对齐不佳，也可能是闭合与抬升时序问题，但记录不足以证明某一个。若检验时序，应从相同的预闭合模拟状态、相同种子和同一策略检查点开始，仅改变规定长度的闭合等待，其余固定，比较接触持续和抬升；本次没有实施这个额外实验。",["evidence/failure_trace.npz","evidence/failure_analysis.json"],80)
    picture(s,"failure_zoom.png",.64,1.72,7.4,4.65)
    txt(s,"Observed",8.25,1.98,4.25,.4,24,BLUE,bold=True)
    txt(s,"Contact lost at step 57\nClose command ≈ +1\nOpening: 0.02240 m\nMaximum can lift: 5.72 mm",8.25,2.55,4.25,1.62,21)
    txt(s,"Hypothesis and test",8.25,4.46,4.25,.4,24,RUST,bold=True)
    txt(s,"Grasp alignment or closure timing.\nTest one timing intervention from the same saved pre-close state.",8.25,5.05,4.25,1.08,20)
    takeaway(s,"The red line is an evidence-based annotation. The proposed intervention has not been executed.")
    s=slide("A complete solution is a chain of evidence","Deliverables and limits", "最后回到学生提交要求。四份完成版 Notebook 保留真实输出，并补齐参考答案。完整证据包包括两组配置、模型、曲线、配对结果、视频和标注；另有两页参考报告。原学生模板和旧输出保留，不要求学生复制相同文字。这个结果只描述本次固定种子的仿真；没有真实机器人验证、没有多个独立训练种子的统计结论，也没有实施新算法验证原因。PPT 有可编辑文字与中文讲解备注，PDF 用于快速预览；教学电脑上的原生 PowerPoint 媒体播放仍需打开确认。",["solution_execution.json","media_provenance.json","verification.json"],45)
    table(s,["Question","Evidence to include","What the answer establishes"],[["Q1","Data contract + demonstration","Observation, action, units and sequence"],["Q2","Configuration + selected checkpoint","How the mapping is learned"],["Q3","Both loss histories + success counts","Prediction loss versus task behavior"],["Q4","Paired rows + initial-position plot","Observed sensitivity; limited coverage"],["Q5","Trace + video + annotated steps","Failure fact, hypothesis and test"]],[1.25,4.85,5.78],row_h=.65,size=19)
    takeaway(s,"One training seed; five cases per condition; simulation only. Retain failures and the limits of inference.")
    name="Manipulation_Lab_Worked_Solutions"
    out=HERE / (name+"_static.pptx" if static else name+".pptx")
    prs.save(out)
    if not static:
        write(HERE / "slides.json", {"slides":records,"total_seconds":sum(r["seconds"] for r in records)})
        (HERE / "speaker_notes_zh.txt").write_text("\n\n".join(f'{r["slide"]:02d} | {r["title"]} | {r["seconds"]} 秒\n{r["notes_zh"]}' for r in records))
    print(out)


if __name__ == "__main__":
    p=argparse.ArgumentParser();p.add_argument("--run",type=Path);p.add_argument("--static",action="store_true");args=p.parse_args()
    d=prepare(args.run) if args.run else read(HERE / "results.json")
    charts(d);build(d,args.static)
