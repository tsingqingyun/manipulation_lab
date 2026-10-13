"""One complete artifact acceptance; never trains or evaluates policies."""
import hashlib
import json
from pathlib import Path
import re
import subprocess
import zipfile
import imageio.v2 as imageio
import imageio_ffmpeg
import nbformat
import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
import pymupdf

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]


def read(p):return json.loads(p.read_text())
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def words(s):return set(re.findall(r"[a-z0-9]+",s.lower()))


def main():
    results=read(HERE/"results.json");run=Path(results["run"])
    original=read(HERE/"evidence/original_execution.json")
    assert original["completed"] and len(original["notebooks"])==4
    notebooks=[]
    for p in sorted((ROOT/"notebooks/solutions_20261013").glob("*.ipynb")):
        n=nbformat.read(p,as_version=4);nbformat.validate(n)
        code=[c for c in n.cells if c.cell_type=="code"]
        assert all(c.execution_count is not None for c in code)
        assert not any(o.output_type=="error" for c in code for o in c.outputs)
        assert "Executed instructor solution" in n.cells[0].source
        notebooks.append({"file":p.name,"code_cells_executed":len(code),"sha256":sha(p)})
    assert len(notebooks)==4
    assert read(HERE/"solution_execution.json")["completed"]
    e=results["experiment"];assert e["complete"] and all(e["checks"].values())
    for v,d in results["variants"].items():
        assert len(d["history"])==d["train"]["epochs"]==50
        assert d["eval"]["errors"]==0 and len(d["eval"]["episodes"])==10
        exp=Path(read(run/"experiment_current.json")["path"])
        ckpt=exp/e["variants"][v]["training"]/"student_policy.pth"
        assert sha(ckpt)==d["eval"]["checkpoint_sha256"]
    trace=np.load(HERE/"evidence/failure_trace.npz")
    marker=results["answers"]["FIRST_ERROR_STEP"]
    assert marker==57 and trace["gripper_contact"][56] and not trace["gripper_contact"][57:].any()
    assert all(f["offset_steps"]==0 for f in results["facts"]["keyframes"])
    for f in results["facts"]["keyframes"]:
        assert trace["video_control_steps"][f["video_frame_index"]]==f["requested_control_step"]
    for entry in read(HERE/"media_provenance.json")["files"]:
        assert sha(HERE/entry["file"])==entry["sha256"]==sha(Path(entry["source"]))
    slides=read(HERE/"slides.json");prs=Presentation(HERE/"Manipulation_Lab_Worked_Solutions.pptx")
    assert len(prs.slides)==len(slides["slides"])==15
    texts=[]
    for i,s in enumerate(prs.slides):
        t=" ".join(sh.text for sh in s.shapes if sh.has_text_frame);texts.append(t)
        assert not re.search(r"[\u4e00-\u9fff]",t)
        assert slides["slides"][i]["notes_zh"] in s.notes_slide.notes_text_frame.text
        for sh in s.shapes:
            assert sh.left>=0 and sh.top>=0
            assert sh.left+sh.width<=prs.slide_width+10 and sh.top+sh.height<=prs.slide_height+10,(i+1,sh.name)
    with zipfile.ZipFile(HERE/"Manipulation_Lab_Worked_Solutions.pptx") as z:
        media=[n for n in z.namelist() if n.endswith(".mp4")];assert len(media)==2
        assert {hashlib.sha256(z.read(n)).hexdigest() for n in media}=={sha(HERE/"assets"/(n+".mp4")) for n in ["success","failure"]}
    pdf=pymupdf.open(HERE/"Manipulation_Lab_Worked_Solutions.pdf");assert len(pdf)==15
    fonts={f[3] for page in pdf for f in page.get_fonts()}
    assert any("SourceSans3" in f for f in fonts) and any("SourceSerif4" in f for f in fonts)
    sheet=Image.new("RGB",(1600,1800),"#FCFBF8")
    for i,page in enumerate(pdf):
        missing=words(texts[i])-words(page.get_text());assert not missing,(i+1,"PDF missing text",missing)
        assert not list(page.annots() or []),"Speaker notes should not become PDF comments"
        dest=HERE/f"preview/slide_{i+1:02d}.png"
        page.get_pixmap(matrix=pymupdf.Matrix(1.25,1.25)).save(dest)
        im=Image.open(dest).convert("RGB");im.thumbnail((533,300));sheet.paste(im,((i%3)*533,(i//3)*360))
        ImageDraw.Draw(sheet).text(((i%3)*533+12,(i//3)*360+310),f"Slide {i+1:02d}",fill="#252C30")
    sheet.save(HERE/"preview/contact_sheet.png")
    report=pymupdf.open(HERE/"Instructor_Solution_Report.pdf");assert len(report)==2
    for i,page in enumerate(report):
        page.get_pixmap(matrix=pymupdf.Matrix(1.35,1.35)).save(HERE/f"preview/report_{i+1:02d}.png")
        assert page.get_text().strip()
    videos=[]
    for name in ["success","failure","demonstration"]:
        p=HERE/"assets"/(name+".mp4")
        subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(),"-v","error","-i",str(p),"-f","null","-"],check=True,capture_output=True)
        with imageio.get_reader(p) as r:
            m=r.get_meta_data();n=r.count_frames()
        videos.append({"file":p.name,"frames":n,"duration_seconds":m["duration"]})
    archive=HERE/"Experiment_Evidence.zip"
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
        m=json.loads(z.read("export_manifest.json"));assert m["answers_complete"]
        for name,h in m["files_sha256"].items():assert hashlib.sha256(z.read(name)).hexdigest()==h
        assert sum(n.endswith("student_policy.pth") for n in z.namelist())==2
        assert json.loads(z.read(m["report"].replace("report.md","submission_manifest.json")))["answers_complete"]
    checks={"passed":True,"notebooks":notebooks,"original_execution_seconds":sum(n["elapsed_seconds"] for n in original["notebooks"]),
        "training_runs":2,"epochs_per_run":50,"evaluation_episodes":20,"evaluation_errors":0,
        "controlled_checks":e["checks"],"slides":15,"speaker_notes_seconds":slides["total_seconds"],
        "pdf_text_coverage":True,"source_fonts_embedded":True,"embedded_video_bytes_match":True,"videos":videos,
        "report_pages":2,"answers_complete":True,"failure_step":marker,"exact_frame_offsets":True,
        "evidence_zip_crc_and_hashes":True,"limits":["Native PowerPoint rendering/playback on the teaching computer is untested.",
        "Browser playback is not part of this acceptance; full MP4 decoding is checked.","One fixed training seed and five paired cases per condition; no strict OOD or real-robot claim.",
        "EGL failed before training; its original attempt is preserved. The completed run uses OSMesa.",
        "Failure cause remains a hypothesis; the proposed intervention has not been executed."]}
    checks["deliverable_sha256"]={p.name:sha(p) for p in [HERE/"Manipulation_Lab_Worked_Solutions.pptx",HERE/"Manipulation_Lab_Worked_Solutions.pdf",HERE/"Instructor_Solution_Report.pdf",archive]}
    (HERE/"verification.json").write_text(json.dumps(checks,indent=2)+"\n")
    print(json.dumps({k:v for k,v in checks.items() if k not in ["notebooks","deliverable_sha256","limits"]},indent=2))


if __name__=="__main__":main()
