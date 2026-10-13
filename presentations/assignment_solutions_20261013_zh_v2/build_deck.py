"""Chinese teaching deck with native diagrams, large charts and verified camera replays."""
import argparse
import json
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
A=HERE/'assets'
D=json.loads((HERE/'results.json').read_text())
OLD_NOTES=json.loads((HERE/'reference_notes.json').read_text())['slides']
BG='FCFBF8'; WHITE='FFFFFF'; INK='252C30'; MUTED='626B70'; BLUE='3F7085'; AMBER='B66C49'; TEAL='6C877D'; RED='B75A50'; LIGHT='E9ECE9'; NAVY='3F7085'
FONT='Lab Classroom Sans'
HEAD='Lab Classroom Serif'
EQUATION=r'\mathcal{L}_{BC}(\theta)=\frac{1}{7N}\sum_{i=1}^{N}\left\|\pi_{\theta}(o_i)-a_i^{*}\right\|_2^2'


def c(x):return RGBColor.from_string(x)


def clean(sh):
    style=sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:sh._element.remove(style)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh


def text(s,t,x,y,w,h,size=22,color=INK,bold=False,align=None,font=FONT):
    b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=b.text_frame
    tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=c(color)
        p.space_after=Pt(0);p.line_spacing=1.12
        if align is not None:p.alignment=align
        for r in p.runs:
            ea=OxmlElement('a:ea');ea.set('typeface',font);rp=r._r.get_or_add_rPr();rp.append(ea);rp.set('lang','zh-CN');rp.set('altLang','zh-CN')
    return b


def box(s,x,y,w,h,fill=WHITE,border=None,rounded=False):
    sh=clean(s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)))
    sh.fill.solid();sh.fill.fore_color.rgb=c(fill)
    if border:sh.line.color.rgb=c(border);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    return sh


def line(s,x1,y1,x2,y2,color=BLUE,width=2,arrow=False):
    sh=clean(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)))
    sh.line.color.rgb=c(color);sh.line.width=Pt(width)
    if arrow:
        end=OxmlElement('a:tailEnd');end.set('type','triangle');sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(end)
    return sh


def circle(s,x,y,d,fill=None,border=BLUE,width=2):
    sh=clean(s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d)))
    sh.line.color.rgb=c(border);sh.line.width=Pt(width)
    if fill:sh.fill.solid();sh.fill.fore_color.rgb=c(fill)
    else:sh.fill.background()
    return sh


def pic(s,name,x,y,w,h,region=(0,0,1,1)):
    p=A/name;iw,ih=Image.open(p).size;l,t,r,b=region;rw=(r-l)*iw;rh=(b-t)*ih
    if rw/rh>w/h:
        keep=rh*w/h/iw;center=(l+r)/2;l,r=center-keep/2,center+keep/2
    else:
        keep=rw*h/w/ih;center=(t+b)/2;t,b=center-keep/2,center+keep/2
    sh=s.shapes.add_picture(str(p),Inches(x),Inches(y),Inches(w),Inches(h));sh.crop_left=l;sh.crop_right=1-r;sh.crop_top=t;sh.crop_bottom=1-b
    return lambda px,py:(x+(px/iw-l)/(r-l)*w,y+(py/ih-t)/(b-t)*h)


def fit(s,name,x,y,w,h):
    iw,ih=Image.open(A/name).size;scale=min(w/iw,h/ih)
    return s.shapes.add_picture(str(A/name),Inches(x+(w-iw*scale)/2),Inches(y+(h-ih*scale)/2),width=Inches(iw*scale),height=Inches(ih*scale))


def red_can(name):
    a=np.asarray(Image.open(A/name).convert('RGB')).astype(float)
    mask=(a[:,:,0]>100)&(a[:,:,0]>1.5*a[:,:,1])&(a[:,:,0]>1.5*a[:,:,2])
    yy,xx=np.where(mask);return float(np.median(xx)),float(np.median(yy))


def tag(s,t,x,y,w,color=BLUE):
    box(s,x,y,w,.39,color,rounded=True);text(s,t,x+.1,y+.045,w-.2,.3,14,WHITE,bold=True)


def note(s,t,dark=False):
    line(s,.65,6.55,12.68,6.55,'D6DBDA',1)
    text(s,t,.7,6.73,11.9,.4,19,BLUE,bold=True)


def fonts():
    # Bundled, distinctly named SC faces keep font selection consistent.
    target=HERE/'fonts/LabClassroomSans-Regular.otf'
    if not target.exists():raise FileNotFoundError('Use the bundled Simplified Chinese fonts')
    font_manager.fontManager.addfont(str(target))
    plt.rcParams.update({'font.family':FONT,'font.size':18,'axes.unicode_minus':False,'figure.facecolor':'#'+WHITE,
        'axes.facecolor':'#'+WHITE,'axes.spines.top':False,'axes.spines.right':False,'axes.labelcolor':'#'+INK,
        'text.color':'#'+INK,'xtick.color':'#'+MUTED,'ytick.color':'#'+MUTED,'axes.edgecolor':'#D7E0EA'})


def charts():
    fonts()
    plt.rcParams['mathtext.fontset']='stix'
    fig=plt.figure(figsize=(12,.9),facecolor='#'+BG)
    fig.text(.02,.14,'$'+EQUATION+'$',fontsize=34,color='#'+INK)
    fig.savefig(A/'bc_equation.png',dpi=350,transparent=True,bbox_inches='tight',pad_inches=.08);plt.close(fig)
    fig,axes=plt.subplots(1,2,figsize=(13.4,4.2))
    for ax,metric,title in zip(axes,['training_loss','validation_loss'],['训练误差','验证误差']):
        for v,col,label in [('lr_a',BLUE,'A · 0.001'),('lr_b',AMBER,'B · 0.0003')]:
            h=D['variants'][v]['history'];xx=[int(r['epoch']) for r in h];yy=[float(r[metric]) for r in h]
            ax.plot(xx,yy,lw=3,color='#'+col,label=label)
        ax.set(title=title,xlabel='训练轮次',ylabel='动作均方误差',xlim=(1,50));ax.grid(axis='y',alpha=.2)
        ax.tick_params(labelsize=18);ax.legend(frameon=False,fontsize=18)
        if metric=='validation_loss':
            ax.axvline(12,color='#'+TEAL,ls='--',lw=2)
            ax.text(13,.047,'第 12 轮\n保存检查点',color='#'+TEAL,fontsize=17)
    fig.tight_layout(w_pad=2);fig.savefig(A/'losses_zh.png',dpi=220);plt.close(fig)
    inspection=D['inspection'];lookup={e['episode']:e['initial_object_position_m'] for e in inspection['episodes']}
    training=np.array([lookup[e] for e in inspection['splits']['train']])
    fig,ax=plt.subplots(figsize=(7,5.0));ax.scatter(training[:,0],training[:,1],s=28,color='#B6C3CF',alpha=.75,label='160 条训练示教起点')
    starts=[]
    for cond,m,filled,label in [('standard','o',True,'标准起点'),('shifted','D',False,'偏移起点')]:
        points=np.array([json.loads(r['initial_object_position']) for r in D['variants']['lr_a']['eval']['episodes'] if r['condition']==cond])
        ax.scatter(points[:,0],points[:,1],s=95,marker=m,facecolors='#'+BLUE if filled else 'none',edgecolors='#'+BLUE,linewidths=2,label=label);starts.append(points)
    for p,q in zip(*starts):ax.plot([p[0],q[0]],[p[1],q[1]],color='#'+BLUE,lw=1.6,alpha=.7)
    ax.set(xlabel='罐子初始 X（米）',ylabel='罐子初始 Y（米）');ax.legend(frameon=False,fontsize=15,loc='upper right')
    ax.set_aspect('equal',adjustable='datalim');ax.tick_params(labelsize=16);fig.tight_layout();fig.savefig(A/'positions_zh.png',dpi=200);plt.close(fig)
    rows=D['variants']['lr_a']['eval']['episodes'];std=[r for r in rows if r['condition']=='standard'];shift=[r for r in rows if r['condition']=='shifted']
    fig,ax=plt.subplots(figsize=(7.5,4.65))
    for i,(r,q) in enumerate(zip(std,shift)):
        x,y=r['steps'],q['steps'];ax.plot([x,y],[i,i],color='#A0B9D3',lw=4,zorder=1)
        if r['stable_success']:
            ax.scatter(x,i,s=100,color='#'+BLUE,zorder=3);ax.scatter(y,i,s=100,marker='D',facecolors='#'+WHITE,edgecolors='#'+BLUE,lw=2,zorder=4)
            ax.annotate(str(x),(x,i),xytext=(-8,13),textcoords='offset points',ha='right',fontsize=14,color='#'+BLUE)
            ax.annotate(str(y),(y,i),xytext=(8,-20),textcoords='offset points',fontsize=14,color='#'+BLUE)
        else:
            ax.scatter(500,i,s=130,marker='x',c='#'+RED,lw=2.5);ax.text(481,i,'均未成功',ha='right',va='center',fontsize=14,color='#'+RED)
    ax.scatter([],[],s=80,color='#'+BLUE,label='标准');ax.scatter([],[],s=80,marker='D',facecolors='#'+WHITE,edgecolors='#'+BLUE,lw=2,label='偏移')
    ax.set(yticks=range(5),yticklabels=[r['source_episode'] for r in std],xlabel='控制步（失败记至 500 步）',xlim=(105,535),xticks=[150,250,350,500],ylim=(4.6,-.6))
    ax.legend(frameon=False,ncol=2,loc='upper center',bbox_to_anchor=(.5,1.14),fontsize=14);ax.grid(axis='x',alpha=.15);ax.tick_params(labelsize=14)
    fig.tight_layout();fig.savefig(A/'paired_steps_zh.png',dpi=200);plt.close(fig)
    trace=np.load(HERE/'evidence/failure_trace.npz');t=np.arange(40,86)
    fig,axes=plt.subplots(3,1,figsize=(8,5.2),sharex=True,gridspec_kw={'height_ratios':[1.4,1,1]})
    for k,label,col in [('eef_position','末端高度',BLUE),('object_position','罐子高度',AMBER)]:axes[0].plot(t,trace[k][40:86,2],lw=2.8,color='#'+col,label=label)
    axes[0].set_ylabel('世界 Z（米）',fontsize=14);axes[0].legend(frameon=False,ncol=2,fontsize=13,loc='upper left')
    axes[1].plot(t,trace['gripper_width'][40:86]*1000,color='#'+BLUE,lw=2.8);axes[1].set_ylabel('开口（毫米）',fontsize=14)
    axes[2].step(t,trace['gripper_contact'][40:86].astype(int),where='post',color='#'+TEAL,lw=2.8);axes[2].set(ylabel='接触',xlabel='控制步（每步 0.05 秒）',yticks=[0,1],yticklabels=['无','有'],ylim=(-.2,1.3))
    for ax in axes:
        ax.axvline(57,color='#'+RED,lw=2,ls='--');ax.set_xlim(40,85);ax.grid(alpha=.12);ax.tick_params(labelsize=13)
    fig.tight_layout(h_pad=.3);fig.savefig(A/'failure_signals_zh.png',dpi=220);plt.close(fig)
    actions=np.asarray(json.loads((HERE/'evidence/demo_0_actions.json').read_text())['actions'])
    fig,ax=plt.subplots(figsize=(12,1.8));ax.step(np.arange(len(actions)),actions[:,-1],where='post',color='#'+BLUE,lw=2.5)
    ax.axvline(52,color='#'+AMBER,ls='--',lw=1.8);ax.text(53,-.5,'第 52 个动作：开始闭合',fontsize=14,color='#'+AMBER)
    ax.set(yticks=[-1,1],yticklabels=['张开 −1','闭合 +1'],xlabel='示教动作索引',xlim=(0,len(actions)-1));ax.tick_params(labelsize=13)
    fig.tight_layout();fig.savefig(A/'demo_action_zh.png',dpi=200);plt.close(fig)


def build(static=False):
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    prs.core_properties.title='机器人操作学习实验｜中文题解'
    prs.core_properties.author='Manipulation Learning Lab'
    records=[]
    corrected_notes={
        5:'本页是人类示教 demo_0 的四个记录状态，依次是四十、五十八、八十二和一百一十六步。它们展示靠近、闭合、搬运与放置的时序。下面横轴是示教动作索引，纵轴是归一化夹爪命令；正一表示闭合，负一表示张开。第一个正命令出现在动作索引五十二，动作与下一状态的对应仍要按记录定义理解。最终照片不能提供这些动作时序，也不能提供每个状态对应的专家标签。这里通过恢复记录状态得到示教画面，与学习策略的闭环运行不同。',
        7:'受控实验的意思是尽量让结果差异只来自一个预先声明的因素。这里两组都训练五十轮，种子十七、批大小二百五十六、网络结构、数据、初始化、设备与评测案例都相同，唯一改变的训练超参数是学习率。代码比较初始权重、数据划分版本和测试清单，四项检查全部通过。两组都是第十二轮验证误差最低，所以保存第十二轮模型。不要把训练预算五十轮误解为必须使用第五十轮模型。A 是测试前已经确定的主要分析对象。',
        8:'两个图分别表示训练集误差和验证集误差，蓝色始终是学习率零点零零一的 A，棕色始终是学习率零点零零零三的 B。横轴是训练轮次，纵轴是对归一化动作计算的均方误差。为了看清各自变化，两幅图的纵轴范围不同，不能用曲线在图上的高度直接跨图比较，必须读数值。两组验证误差均在第十二轮达到最低，而训练误差继续下降。第五十轮 A 的训练误差更低，只说明对训练示例拟合更紧，不代表它在闭环里一定更好。',
        10:'灰点是训练示教的起始位置，每条轨迹一个点，不是整条运动路径。蓝色实心圆是标准测试起点，空心菱形是偏移后的起点，细线连接同一个配对案例。右侧只是世界 X 方向增加四厘米的概念示意。源状态和随机种子保持相同，偏移后的罐子仍需处在源料箱内部。这种扰动可以检查位置敏感性，但二维位置图省略了朝向、手臂与夹爪状态等输入，偏移后也可能落在训练位置覆盖内，因此不能当作严格的分布外泛化证明。',
        11:'左图保留主要策略 A 的全部五个配对案例。圆点表示标准条件，空心菱形表示偏移条件，横轴是控制步。成功案例的数字是达到稳定放置的时刻；两个叉号在五百步位置，表示到了上限仍未成功，不能解释成较慢完成。所有配对的成败标签都没有改变，但 demo_185 从一百四十一步变为一百五十六步，增加十五步即零点七五秒。总成功率相同不等于运动过程完全相同。抓取几何可能是解释之一，但单看这张时间图还不能证明机制。',
        14:'左侧三个坐标轴依次显示末端和罐子的世界高度、以毫米显示的夹爪开口、以及有无接触。红虚线都指向第五十七步，是人工根据原始记录标注的接触丢失，不是系统自动认定的根因。末端随后明显上升，罐子最高只比初始高五点七二毫米。右侧把事实、原因假设和检验方案分开：可疑因素包括对齐和闭合抬升时序；要检验时序，可从完全相同的预闭合仿真与控制器状态开始，只改变固定五步的闭合等待，然后恢复原策略。比较接触持续、罐子抬高和最终放置。本次没有执行这个干预，因此不能说已经修复。'
    }
    def slide(title,section,old=None,notes=None,seconds=60,dark=False):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=c(BG)
        n=len(prs.slides);text(s,section,.66,.36,10.9,.32,13,BLUE,bold=True)
        if title:text(s,title,.63,.98,12.08,.85,32,INK,font=HEAD)
        text(s,'机器人操作学习 · 中文题解',.66,7.24,9,.18,9,MUTED)
        text(s,f'{n:02d} / 15',11.86,7.18,.82,.25,11,MUTED,align=PP_ALIGN.RIGHT)
        narration=notes or corrected_notes.get(old) or OLD_NOTES[old-1]['notes_zh'];sources=OLD_NOTES[old-1]['sources'] if old else ['results.json','replay_provenance.json']
        if n==2:sources=['assets/success_hd.mp4','replay_provenance.json','evidence/evaluation_config.json']
        if n==5:sources=['replay_provenance.json','evidence/demo_0_actions.json']
        if n==15:sources=['solution_execution.json','replay_provenance.json','verification.json']
        s.notes_slide.notes_text_frame.text=f'第 {n} 页｜{seconds} 秒\n{narration}\n\n来源：\n'+'\n'.join(sources)
        records.append({'slide':n,'title':title or '从示教学动作，到解释机器人为什么失败','notes_zh':narration,'seconds':seconds,'sources':sources})
        return s
    def movie(s,label,x,y,w):
        h=w*544/960
        if static:pic(s,label+'_053.png' if label=='failure' else 'success_121.png',x,y,w,h)
        else:clean(s.shapes.add_movie(str(A/(label+'_hd.mp4')),Inches(x),Inches(y),Inches(w),Inches(h),poster_frame_image=str(A/(label+'_053.png' if label=='failure' else 'success_121.png')),mime_type='video/mp4'))
    s=slide('', '机器人操作学习实验 · 参考题解',old=1,seconds=35,dark=True)
    text(s,'从示教学动作，\n到解释机器人\n为什么失败',.72,1.24,6.4,2.9,40,INK,font=HEAD)
    text(s,'读数据 → 训练策略 → 配对评测 → 定位失败',.77,4.51,6.4,.5,19,MUTED)
    for i,(num,label) in enumerate([('4','份完整 Notebook'),('2','种学习率'),('20','次实际评测')]):
        text(s,num,.78+i*2.0,5.34,1.55,.63,36,BLUE,bold=True);text(s,label,.8+i*2.0,6.05,1.8,.34,14,MUTED)
    pic(s,'detail_065.png',7.28,1.07,5.42,5.5,region=(.08,0,.9,1))
    tag(s,'真实轨迹：夹爪上升，罐子留下',7.55,5.87,4.78,RED)

    s=slide('任务完成，要看“松开以后”','先看结果，再解释学习过程',old=2)
    movie(s,'success',.68,1.98,7.6);tag(s,'点击视频播放 · 按原动作半速重放',.88,5.8,5.55,BLUE)
    for i,(title,desc) in enumerate([('抓起','夹爪带动物体离开桌面'),('放入','罐子进入目标区域并留有余量'),('松开','夹爪张开，而且不再接触罐子')]):
        circle(s,8.7,2.03+i*.91,.35,BLUE,BLUE);text(s,str(i+1),8.75,2.04+i*.91,.25,.26,14,WHITE,bold=True)
        text(s,title,9.23,2.02+i*.91,3.3,.36,22,bold=True);text(s,desc,9.23,2.45+i*.91,3.25,.4,16,MUTED)
    text(s,'再稳定 1 秒',8.75,5.04,3.6,.56,30,TEAL,bold=True)
    text(s,'位移 ≤ 1 cm · 速度 ≤ 5 cm/s',8.76,5.76,3.9,.32,15,MUTED)
    note(s,'这 1 秒里策略继续执行；抬起罐子只是任务的一步。')

    s=slide('Q1｜机器人“看到”的，是 16 个状态值','观察量：决策时已知的信息',old=3)
    project=pic(s,'detail_053.png',.67,1.92,6.0,4.13,region=(.06,.08,.94,.95))
    px,py=project(*red_can('detail_053.png'));circle(s,px-.46,py-.46,.92,None,AMBER,2.7)
    tag(s,'罐子的世界位姿',.89,5.46,2.75,AMBER)
    line(s,3.38,5.48,px,py+.42,AMBER,2)
    for i,(n,title,desc,col) in enumerate([('3','末端位置','世界 XYZ · 米',BLUE),('4','末端朝向','四元数 · xyzw',BLUE),('2','夹爪关节位置','两个手指的位置 · 米',TEAL),('7','罐子位姿','位置 3 + 朝向 4',AMBER)]):
        y=1.94+i*1.02;box(s,7.01,y,5.65,.89,WHITE,rounded=True)
        text(s,n,7.24,y+.11,.7,.6,32,col,bold=True);text(s,title,8.12,y+.1,4.2,.34,20,bold=True);text(s,desc,8.13,y+.52,4.2,.3,15,MUTED)
    note(s,'3 + 4 + 2 + 7 = 16；训练与执行都使用 object[7:14] 的世界位姿。')

    s=slide('Q1｜7 维动作，是下一步的控制命令','动作：归一化命令与实际位移要分清',old=4)
    for i,label in enumerate(['Δx','Δy','Δz','Δrx','Δry','Δrz','开 / 合']):
        x=.75+i*1.79;col=BLUE if i<3 else (TEAL if i<6 else AMBER)
        box(s,x,2.0,1.64,.9,col,rounded=True);text(s,label,x+.05,2.18,1.54,.53,27,WHITE,bold=True,align=PP_ALIGN.CENTER)
    text(s,'平移增量 × 3',.92,3.1,4.55,.4,23,BLUE,bold=True);text(s,'世界坐标 · 上限 ±0.05 m',.92,3.62,4.7,.4,18,MUTED)
    text(s,'旋转增量 × 3',6.08,3.1,4.55,.4,23,TEAL,bold=True);text(s,'轴角表示 · 上限 ±0.5 rad',6.08,3.62,4.6,.4,18,MUTED)
    box(s,.77,4.52,6.22,1.5,WHITE,rounded=True)
    text(s,'平移命令 0.2',1.02,4.87,2.46,.5,25,bold=True);line(s,3.55,5.12,4.2,5.12,BLUE,3,True);text(s,'请求 1 cm',4.45,4.87,2.37,.5,26,BLUE,bold=True)
    text(s,'实际位移还受动力学与接触影响',1.03,5.57,5.58,.31,16,MUTED)
    box(s,7.25,4.52,5.37,1.5,WHITE,rounded=True)
    box(s,7.59,4.8,.14,.88,INK);box(s,8.89,4.8,.14,.88,INK);line(s,7.85,5.2,8.2,5.2,AMBER,3,True);line(s,8.79,5.2,8.45,5.2,AMBER,3,True)
    text(s,'两个关节\n一个开合命令',9.28,4.73,3.1,.78,20,bold=True);text(s,'正值闭合 · 负值张开',9.28,5.64,3.1,.3,15,MUTED)
    note(s,'控制频率 20 Hz：每步 0.05 秒；夹爪命令无量纲，测得的开口用米表示。')

    s=slide('Q1｜轨迹告诉我们“先做什么，再做什么”','人类示教 demo_0 · 118 个记录步',old=5,seconds=45)
    for i,(step,label) in enumerate([(40,'靠近'),(58,'闭合'),(82,'搬运'),(116,'放置')]):
        x=.68+i*3.16;pic(s,f'demo_{step:03d}.png',x,1.98,3.01,1.93)
        tag(s,f'{label} · 状态 {step}',x+.09,3.46,2.73,BLUE)
        if i<3:line(s,x+3.04,2.81,x+3.13,2.81,BLUE,2,True)
    fit(s,'demo_action_zh.png',.78,4.3,11.8,1.78)
    note(s,'最终照片没有动作时序；示教轨迹提供“当前观察 → 专家动作”的训练样本。')

    s=slide('Q2｜行为克隆：根据状态，预测专家动作','行为克隆（BC）：把示教变成监督学习',old=6)
    for x,w,title,desc,col in [(.74,3.08,'输入：16 维观察','只用训练集统计量归一化',BLUE),(4.33,3.73,'策略网络','三层隐藏层：256 → 256 → 256',TEAL),(8.57,4.01,'输出：7 维动作','与同一时刻的专家动作比较',AMBER)]:
        box(s,x,2.06,w,1.45,WHITE,rounded=True);text(s,title,x+.19,2.29,w-.36,.4,24,col,bold=True);text(s,desc,x+.19,2.93,w-.36,.38,15,MUTED)
    line(s,3.9,2.75,4.24,2.75,BLUE,2.3,True);line(s,8.11,2.75,8.48,2.75,TEAL,2.3,True)
    fit(s,'bc_equation.png',.88,3.78,11.63,.85)
    text(s,'o：观察　a*：专家动作　πθ：策略　N：样本数；对 7 个动作分量取均值。',.91,4.89,11.6,.39,18,MUTED)
    x=.88
    for width,label,col in [(8.68,'训练 160 条：学习权重与归一化',BLUE),(1.42,'验证 20 条',TEAL),(1.42,'测试 20 条',AMBER)]:
        box(s,x,5.47,width,.61,col);text(s,label,x+.1,5.6,width-.2,.34,15,WHITE,bold=True,align=PP_ALIGN.CENTER);x+=width+.04
    note(s,'整条轨迹进入同一个集合；验证集选检查点，测试集只用于评测。')

    s=slide('Q2｜只改学习率，其他条件保持一致','受控实验：变化来自哪一个因素？',old=7)
    box(s,.73,1.98,11.9,.74,NAVY,rounded=True);text(s,'固定：相同数据与初始权重  /  种子 17  /  批大小 256  /  相同网络与设备',.95,2.19,11.4,.38,19,WHITE)
    for i,(label,lr,col) in enumerate([('A · lr_a','0.001',BLUE),('B · lr_b','0.0003',AMBER)]):
        x=.74+i*6.17;box(s,x,3.03,5.74,2.36,WHITE,rounded=True);tag(s,label,x+.22,3.24,1.77,col)
        text(s,'学习率',x+.25,3.98,1.44,.4,20,MUTED);text(s,lr,x+1.79,3.66,3.59,.95,48,col,bold=True)
        text(s,'训练 50 轮 → 保存第 12 轮',x+.25,4.79,5.2,.4,22,bold=True)
    text(s,'相同 5 个标准案例 + 5 个位置偏移案例；4 项一致性检查全部通过。',.9,5.76,11.8,.4,20,TEAL,bold=True)
    note(s,'A 在测试前已固定为主要分析模型；检查点按验证误差选择。')

    s=slide('Q3｜训练还在变好，验证误差却不再下降','读曲线：蓝色 A，棕色 B；两个图的纵轴范围各自标明',old=8)
    box(s,.72,1.88,11.9,3.85,WHITE,rounded=True);fit(s,'losses_zh.png',.76,1.91,11.82,3.78)
    text(s,'第 50 轮训练误差',.89,5.9,3.32,.35,18,MUTED);text(s,'A 0.01772  <  B 0.02092',4.1,5.83,7.5,.48,26,BLUE,bold=True)
    note(s,'最低验证误差均在第 12 轮；继续降低训练误差，不代表新状态上的表现改善。')

    s=slide('Q3｜训练误差更低，放置成功次数却更少','任务表现：每一颗圆点代表一次真实评测',old=9,dark=True)
    text(s,'标准条件',4.1,2.0,3.45,.42,23,INK,bold=True);text(s,'位置偏移',8.72,2.0,3.45,.42,23,INK,bold=True)
    for j,(v,label,col) in enumerate([('lr_a','A · 0.001',BLUE),('lr_b','B · 0.0003',AMBER)]):
        y=2.98+j*1.55;text(s,label,.79,y+.05,2.78,.55,27,BLUE if j==0 else AMBER,bold=True)
        for i,cond in enumerate(['standard','shifted']):
            rows=[r for r in D['variants'][v]['eval']['episodes'] if r['condition']==cond];x=4.0+i*4.64
            for k,r in enumerate(rows):
                circle(s,x+k*.5,y,.34,TEAL if r['stable_success'] else None,TEAL if r['stable_success'] else MUTED,2)
                if not r['stable_success']:line(s,x+k*.5+.05,y+.05,x+k*.5+.29,y+.29,MUTED,1.4)
            n=sum(r['stable_success'] for r in rows);text(s,f'{n}/5',x+2.72,y-.19,1.47,.7,40,INK,bold=True)
    text(s,'评测使用第 12 轮检查点，不能用第 50 轮训练误差替代任务结果。',.86,5.95,11.8,.39,19,MUTED)
    note(s,'每种条件只有 5 个案例、1 个训练种子；这里展示观察结果，不能据此给学习率做普遍排名。',True)

    s=slide('Q4｜位置偏移测试，只把罐子起点平移 4 cm','先认清图上的点和线，再讨论泛化',old=10)
    box(s,.7,1.92,7.04,4.35,WHITE,rounded=True);fit(s,'positions_zh.png',.76,1.95,6.9,4.24)
    circle(s,8.57,2.42,.24,BLUE,BLUE);circle(s,11.17,2.42,.24,None,BLUE);line(s,8.99,2.54,10.91,2.54,BLUE,3,True)
    text(s,'标准',8.34,2.98,1.6,.4,20);text(s,'偏移',10.93,2.98,1.6,.4,20)
    text(s,'世界 X +0.04 m',8.42,3.69,4.04,.6,28,BLUE,bold=True)
    text(s,'同一源状态 · 同一种子\n偏移后仍在源料箱内',8.42,4.59,4.02,1.0,21)
    text(s,'点 = 初始位置\n线 = 两个配对起点的对应关系',8.42,5.59,4.1,.65,16,MUTED)
    note(s,'二维位置图既不是运动轨迹，也不能证明完整观察空间里的分布外（OOD）泛化。')

    s=slide('Q4｜成败标签相同，完成时间仍会变化','主要模型 A：五个案例全部保留',old=11)
    box(s,.7,2.03,7.53,4.24,WHITE,rounded=True);fit(s,'paired_steps_zh.png',.71,2.04,7.5,4.18)
    tag(s,'同一案例：demo_185 / 种子 102',8.55,2.05,4.07,BLUE)
    text(s,'141 → 156',8.62,2.87,4.1,.88,41,BLUE,bold=True)
    text(s,'控制步',8.7,3.76,3.9,.4,19,MUTED)
    text(s,'+0.75 秒',8.64,4.4,4.0,.74,35,bold=True)
    text(s,'两种条件都成功\n但执行过程并不完全相同',8.65,5.48,3.9,.77,20)
    note(s,'500 步标记的是失败上限，不是成功时间；时间变化本身也不能证明某个原因。')

    s=slide('Q5｜第 57 步：夹爪失去接触，随后空手抬起','原始失败 standard_00 · 主要模型 A',notes='先播放原失败前一百零一步的高清重放，半速便于看清抓持过程。重放沿用已保存动作，没有重新调用策略；逐步比较位置、开口和接触等记录，全部与原轨迹一致。关键不在最终五百步没有完成，而在四十九至五十六步尚有接触，五十七步首次失去接触，随后罐子没有随夹爪持续上升。标注五十七步为最早有直接证据支持的接触丢失，不将其等同于最早因果错误。',seconds=70)
    movie(s,'failure',.7,1.99,7.6);tag(s,'点击播放 · 前 101 步 · 半速重放',.91,5.85,5.55,RED)
    text(s,'49–56 步',8.73,2.2,3.8,.62,31,TEAL,bold=True);text(s,'夹爪与罐子有接触',8.78,2.94,3.8,.4,21)
    line(s,9.12,3.5,9.12,4.05,RED,3,True)
    text(s,'第 57 步',8.72,4.24,3.8,.62,34,RED,bold=True);text(s,'接触变为“无”\n闭合命令仍约为 +1',8.77,5.03,3.8,.9,22)
    note(s,'这是可观察的抓持丢失；“对齐不佳”或“闭合过晚”仍需要额外证据。')

    s=slide('Q5｜把错误时刻放大，和状态数值一一对应','同一失败轨迹 · 更近视角 · 与原记录逐步比对一致',notes='这四张图片不是放大低分辨率原视频，而是在同一个仿真初态下重放已保存动作，用更近的固定相机渲染。四个状态分别是四十五、五十三、五十七和六十五步。五十三步有接触；五十七步失去接触；六十五步夹爪几乎合拢并上升，而罐子仍留在桌面附近。图中开口为了易读转换成毫米，原始记录单位是米。红色强调的是原轨迹第57步，新的视角不改变证据的来源，也不能凭一张图断言确切根因。',seconds=65)
    tr=np.load(HERE/'evidence/failure_trace.npz')
    for i,(t,label,col) in enumerate([(45,'接近',BLUE),(53,'发生接触',TEAL),(57,'接触断开',RED),(65,'空手上升',RED)]):
        x=.68+i*3.16;box(s,x,1.98,3.0,4.31,WHITE,rounded=True)
        project=pic(s,f'detail_{t:03d}.png',x+.02,2.06,2.96,2.57,region=(.19,.1,.77,.98))
        px,py=project(*red_can(f'detail_{t:03d}.png'));circle(s,px-.31,py-.31,.62,None,col,2)
        tag(s,f'{t} 步 · {label}',x+.15,4.71,2.7,col)
        text(s,f'开口 {tr["gripper_width"][t]*1000:.1f} mm',x+.18,5.33,2.7,.4,21,bold=True)
        text(s,f'罐子 Z = {tr["object_position"][t,2]:.4f} m',x+.18,5.91,2.7,.3,15,MUTED)
    note(s,'动作 a[t−1] 导致状态 s[t]；这些图片标的是控制步，不是视频帧号。')

    s=slide('Q5｜事实、假设、检验，分三步说清楚','诊断：先找最早异常，再讨论可能原因',old=14,seconds=80)
    box(s,.68,1.94,7.6,4.36,WHITE,rounded=True);fit(s,'failure_signals_zh.png',.72,1.97,7.51,4.27)
    for y,num,title,desc,col in [(2.0,'1','事实','57 步失去接触\n罐子最大只抬高 5.72 mm',RED),(3.46,'2','假设','抓取对齐不佳？\n闭合与抬升时序不当？',AMBER),(4.92,'3','检验','从相同预闭合状态开始\n只改变固定 5 步的闭合等待',TEAL)]:
        circle(s,8.65,y,.34,col,col);text(s,num,8.73,y+.01,.2,.3,14,WHITE,bold=True)
        text(s,title,9.19,y-.04,3.3,.43,24,col,bold=True);text(s,desc,8.68,y+.55,3.96,.78,18)
    note(s,'红虚线是人工证据标注；右侧的改进实验尚未执行，不能宣称已经修复。')

    s=slide('一份好题解，让每个结论都能找到证据','回到 Q1–Q5：看懂数据、理解学习、解释结果',old=15,seconds=45,dark=True)
    for i,(num,title,desc,col) in enumerate([('01','观察与动作','16 维状态 → 7 维命令\n坐标、单位、时序说清楚',BLUE),('02','训练与选择','同样条件，只改学习率\n验证集选择第 12 轮',TEAL),('03','评测与泛化','成功次数 + 五个配对案例\n位置偏移不等于严格 OOD',AMBER),('04','失败与原因','第 57 步是观察到的事实\n原因假设需要受控检验',RED)]):
        x=.74+(i%2)*6.18;y=2.03+(i//2)*2.05;box(s,x,y,5.74,1.8,LIGHT,rounded=True)
        text(s,num,x+.22,y+.23,.9,.64,31,BLUE,bold=True);text(s,title,x+1.37,y+.23,4.1,.45,24,INK,bold=True)
        text(s,desc,x+1.39,y+.89,4.1,.74,18,MUTED)
    note(s,'四份 Notebook 已完整执行；本课讨论的是一次固定种子的仿真证据。',True)
    name='Manipulation_Lab_Solutions_ZH'
    target=HERE/(name+'_static.pptx' if static else name+'.pptx');prs.save(target)
    if not static:
        (HERE/'slides.json').write_text(json.dumps({'slides':records,'seconds':sum(r['seconds'] for r in records),'visible_language':'Chinese explanations; English mathematical notation','notes_language':'Chinese','equations':[{'latex':EQUATION,'asset':'assets/bc_equation.png'}],'style':'Original cream background, blue-gray and warm-brown palette, serif headings'},ensure_ascii=False,indent=2)+'\n')
        (HERE/'speaker_notes_zh.txt').write_text('\n\n'.join(f'{r["slide"]:02d}｜{r["title"]}｜{r["seconds"]} 秒\n{r["notes_zh"]}' for r in records)+'\n')
    print(target)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--static',action='store_true');p.add_argument('--skip-charts',action='store_true');args=p.parse_args()
    if not args.skip_charts:charts()
    build(args.static)
